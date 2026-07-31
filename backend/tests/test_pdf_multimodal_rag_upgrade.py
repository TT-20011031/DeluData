"""
PDF 多模态 RAG 升级测试
"""
import asyncio
import io
import sys
from contextlib import asynccontextmanager
from pathlib import Path
from urllib.parse import urlparse, parse_qs
from unittest.mock import MagicMock

import fitz
import pytest
from fastapi import HTTPException

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.api.knowledge.images import get_image_meta
import app.api.knowledge.images as knowledge_images_module
from app.core.llm.async_llm import AsyncLLMClient
from app.core.llm.async_embedding import AsyncEmbeddingClient
from app.core.rag.document_parser import DocumentParser
from app.core.rag.document_ingestor import DocumentIngestor
from app.core.utils.image_service import ImageService
from app.core.security.auth import User
from app.config import get_settings
from app.supervisor.nodes.synthesizer import (
    _collect_mm_image_urls,
    _looks_like_html_payload,
    _safe_non_negative_int,
)
import app.supervisor.nodes.synthesizer as synthesizer_module
import app.supervisor.nodes.direct_execute as direct_execute_module
import app.core.rag.document_ingestor as document_ingestor_module
from app.tools.doc_tool import (
    _build_citation_file_slot,
    _parse_page_numbers,
    select_chunks_for_synthesizer,
    select_page_images_for_synthesizer,
)


class DummyChunk:
    def __init__(
        self,
        *,
        content: str,
        score: float,
        source_file: str = "demo.pdf",
        file_id: str = "file_1",
        page: int = 1,
        chunk_id: str = "",
    ):
        self.content = content
        self.score = score
        self.rerank_score = score
        self.source_file = source_file
        self.page_number = page
        self.chunk_id = chunk_id or f"chunk_{page}_{score}"
        self.metadata = {
            "file_id": file_id,
            "page_numbers": str(page),
            "header_path": f"h-{page}-{score}",
            "type": "text",
        }


def test_document_parser_watermark_clean_and_ocr_trigger():
    parser = DocumentParser(max_file_size=1024 * 1024)

    cleaned = parser._clean_page_text("标准分享网 免费下载 Hello")
    assert "标准分享网" not in cleaned
    assert "免费下载" not in cleaned

    assert parser._should_run_ocr(extract_failed=True, cleaned_text_length=999)
    assert parser._should_run_ocr(extract_failed=False, cleaned_text_length=49)
    assert not parser._should_run_ocr(extract_failed=False, cleaned_text_length=50)
    assert parser._should_run_ocr(
        extract_failed=False,
        cleaned_text_length=500,
        image_area_ratio=0.31,
    )


def test_document_parser_parse_pdf_detailed_uses_ocr_for_short_text(tmp_path):
    pdf_path = tmp_path / "short_text.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "abc")
    doc.save(str(pdf_path))
    doc.close()

    parser = DocumentParser(max_file_size=10 * 1024 * 1024)
    parser._ocr_from_image_bytes = lambda png_bytes: ("ocr text", None)  # type: ignore

    result = parser._parse_pdf_detailed_sync(pdf_path)
    assert result["total_pages"] == 1
    assert result["pages"][0]["used_ocr"] is True
    assert "ocr text" in result["text"]


def test_document_parser_parse_pdf_detailed_ocr_large_image_page_keeps_native_text(tmp_path):
    from PIL import Image

    pdf_path = tmp_path / "mixed_image_text.pdf"
    image = Image.new("RGB", (1200, 800), "white")
    image_bytes = io.BytesIO()
    image.save(image_bytes, format="PNG")
    image_bytes.seek(0)

    doc = fitz.open()
    page = doc.new_page(width=600, height=800)
    page.insert_textbox(
        fitz.Rect(40, 40, 560, 180),
        "native selectable text " * 20,
        fontsize=12,
    )
    page.insert_image(fitz.Rect(40, 260, 560, 760), stream=image_bytes.getvalue())
    doc.save(str(pdf_path))
    doc.close()

    parser = DocumentParser(max_file_size=10 * 1024 * 1024)
    parser._ocr_from_image_bytes = lambda png_bytes: ("control panel cleaning standard", None)  # type: ignore

    result = parser._parse_pdf_detailed_sync(pdf_path)

    assert result["total_pages"] == 1
    assert result["pages"][0]["used_ocr"] is True
    assert result["pages"][0]["image_area_ratio"] >= 0.30
    assert "native selectable text" in result["text"]
    assert "control panel cleaning standard" in result["text"]


def test_document_parser_parse_docx_with_page_break_generates_page_markers():
    from docx import Document
    from docx.enum.text import WD_BREAK
    from uuid import uuid4

    tmp_dir = BACKEND_ROOT / "tmp_test_docx_page_markers"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    docx_path = tmp_dir / f"with_page_break_{uuid4().hex}.docx"
    doc = Document()
    doc.add_paragraph("第一页内容")
    run = doc.add_paragraph().add_run()
    run.add_break(WD_BREAK.PAGE)
    doc.add_paragraph("第二页内容")
    doc.save(str(docx_path))

    try:
        parser = DocumentParser(max_file_size=10 * 1024 * 1024)
        text = parser._parse_docx(docx_path)

        assert "[PAGE:1]" in text
        assert "[PAGE:2]" in text
        assert text.index("[PAGE:1]") < text.index("第一页内容")
        assert text.index("[PAGE:2]") < text.index("第二页内容")
    finally:
        try:
            docx_path.unlink(missing_ok=True)
        except Exception:
            pass


def test_document_parser_parse_docx_with_multiple_page_breaks_generates_incremental_markers():
    from docx import Document
    from docx.enum.text import WD_BREAK
    from uuid import uuid4

    tmp_dir = BACKEND_ROOT / "tmp_test_docx_page_markers"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    docx_path = tmp_dir / f"with_multi_page_break_{uuid4().hex}.docx"
    doc = Document()
    doc.add_paragraph("第一页")
    run = doc.add_paragraph().add_run()
    run.add_break(WD_BREAK.PAGE)
    run.add_break(WD_BREAK.PAGE)
    doc.add_paragraph("第三页")
    doc.save(str(docx_path))

    try:
        parser = DocumentParser(max_file_size=10 * 1024 * 1024)
        text = parser._parse_docx(docx_path)

        assert "[PAGE:1]" in text
        assert "[PAGE:2]" in text
        assert "[PAGE:3]" in text
        assert text.index("[PAGE:3]") < text.index("第三页")
    finally:
        try:
            docx_path.unlink(missing_ok=True)
        except Exception:
            pass


def test_document_parser_ocr_from_image_bytes_tiling_and_dedupe(tmp_path):
    pdf_path = tmp_path / "tall_page.pdf"
    doc = fitz.open()
    doc.new_page(width=800, height=5000)
    doc.save(str(pdf_path))
    doc.close()

    doc = fitz.open(str(pdf_path))
    page = doc[0]
    pixmap = page.get_pixmap(dpi=72, alpha=False)
    png_bytes = pixmap.tobytes("png")
    doc.close()

    parser = DocumentParser(max_file_size=20 * 1024 * 1024)

    calls = {"count": 0}

    def fake_engine(_img, **_kwargs):
        calls["count"] += 1
        idx = calls["count"]
        box = [[0, 10], [20, 10], [20, 30], [0, 30]]
        return ([[box, f"tile-{idx}"], [box, f"tile-{idx}"]], None)

    parser._get_ocr_engine = lambda: fake_engine  # type: ignore
    text, err = parser._ocr_from_image_bytes(png_bytes)

    assert err is None
    assert calls["count"] >= 2
    lines = [line for line in text.splitlines() if line.strip()]
    assert len(lines) == calls["count"]
    assert lines[0] == "tile-1"


def test_document_parser_ocr_from_image_bytes_rapidocr_output_compat(tmp_path):
    pdf_path = tmp_path / "rapidocr_output.pdf"
    doc = fitz.open()
    doc.new_page(width=800, height=1200)
    doc.save(str(pdf_path))
    doc.close()

    doc = fitz.open(str(pdf_path))
    page = doc[0]
    pixmap = page.get_pixmap(dpi=72, alpha=False)
    png_bytes = pixmap.tobytes("png")
    doc.close()

    parser = DocumentParser(max_file_size=20 * 1024 * 1024)

    class FakeRapidOutput:
        def __init__(self):
            import numpy as np
            self.boxes = np.array(
                [
                    [[0, 10], [20, 10], [20, 30], [0, 30]],
                    [[0, 40], [20, 40], [20, 60], [0, 60]],
                ]
            )
            self.txts = ["line-1", "line-2"]

    def fake_engine(_img, **kwargs):
        if "det_db_box_thresh" in kwargs:
            raise TypeError("unexpected keyword")
        return FakeRapidOutput()

    parser._get_ocr_engine = lambda: fake_engine  # type: ignore
    text, err = parser._ocr_from_image_bytes(png_bytes)

    assert err is None
    assert "line-1" in text
    assert "line-2" in text


def test_select_chunks_for_synthesizer_respects_6_to_8():
    chunks = [
        DummyChunk(content=f"content {i} " * 10, score=1.0 - i * 0.01, page=i + 1, chunk_id=f"c{i}")
        for i in range(12)
    ]

    selected, used_tokens, _ = select_chunks_for_synthesizer(
        chunks=chunks,
        token_budget=20000,
        min_chunks=6,
        max_chunks=8,
    )
    assert 6 <= len(selected) <= 8
    assert used_tokens > 0


def test_select_page_images_top3_window_dedupe_and_cap():
    chunks = [
        DummyChunk(content="a", score=0.95, file_id="pdf_1", page=10, chunk_id="a"),
        DummyChunk(content="b", score=0.90, file_id="pdf_1", page=10, chunk_id="b"),
        DummyChunk(content="c", score=0.85, file_id="pdf_1", page=11, chunk_id="c"),
        DummyChunk(content="d", score=0.80, file_id="pdf_2", page=3, chunk_id="d"),
    ]
    file_info_map = {
        "pdf_1": {"type": "pdf", "storage_path": "x.pdf"},
        "pdf_2": {"type": "pdf", "storage_path": "y.pdf"},
    }

    selected = select_page_images_for_synthesizer(
        chunks=chunks,
        file_info_map=file_info_map,
        anchor_chunks=3,
        page_window=1,
        max_images=9,
    )
    keys = {(item["file_id"], item["page_number"]) for item in selected}
    assert len(selected) == len(keys)

    clipped = select_page_images_for_synthesizer(
        chunks=chunks,
        file_info_map=file_info_map,
        anchor_chunks=3,
        page_window=1,
        max_images=3,
    )
    assert len(clipped) == 3
    assert clipped[0]["file_id"] == "pdf_1"
    assert clipped[0]["page_number"] == 10


def test_collect_mm_image_urls_dedupe():
    results = [
        {
            "meta": {
                "mm_evidence": {
                    "images": [
                        {"file_id": "f1", "page_number": 1, "url": "https://a/1"},
                        {"file_id": "f1", "page_number": 1, "url": "https://a/1_dup"},
                        {"file_id": "f1", "page_number": 2, "url": "https://a/2"},
                    ]
                }
            }
        },
        {
            "meta": {
                "mm_evidence": {
                    "images": [
                        {"file_id": "f2", "page_number": 1, "url": "https://b/1"},
                    ]
                }
            }
        },
    ]

    urls = _collect_mm_image_urls(results, max_images=2)
    assert urls == ["https://a/1", "https://a/2"]


def test_safe_non_negative_int_and_html_payload_helpers():
    assert _safe_non_negative_int("7", default=9) == 7
    assert _safe_non_negative_int("-3", default=9) == 0
    assert _safe_non_negative_int("not-a-number", default=9) == 9

    assert _looks_like_html_payload("<!DOCTYPE html><html><body>x</body></html>")
    assert _looks_like_html_payload("   <html><body>x</body></html>")
    assert not _looks_like_html_payload("plain text only")
    assert not _looks_like_html_payload(None)


def test_synthesizer_multimodal_failure_fallback_to_text_only(monkeypatch):
    class FakeLLM:
        def __init__(self):
            self.calls = []

        async def chat_stream(self, messages, model=None):
            self.calls.append(messages)
            user_content = messages[1]["content"]
            if isinstance(user_content, list):
                raise RuntimeError("multimodal fetch failed")
            for chunk in ("fallback ", "answer"):
                yield chunk

    fake_llm = FakeLLM()
    monkeypatch.setattr(synthesizer_module, "get_async_llm", lambda: fake_llm)

    state = {
        "intent_type": "tool_use",
        "user_query": "请总结",
        "execution_results": [
            {
                "step_id": "doc_1",
                "worker": "doc_worker",
                "result": "这是结果",
                "meta": {
                    "mm_evidence": {
                        "images": [
                            {
                                "file_id": "file_1",
                                "page_number": 1,
                                "url": "https://img.example.com/1.png",
                            }
                        ]
                    }
                },
            }
        ],
        "task_plan": [],
        "messages": [],
        "pending_artifacts": {},
    }

    output = asyncio.run(synthesizer_module.synthesizer_node(state))

    assert output["final_answer"] == "fallback answer"
    assert len(fake_llm.calls) == 2
    assert isinstance(fake_llm.calls[0][1]["content"], list)
    assert isinstance(fake_llm.calls[1][1]["content"], str)


def test_async_llm_multimodal_normalization():
    client = AsyncLLMClient()
    messages = [
        {
            "role": "user",
            "content": [
                {"text": "问题"},
                {"image": "https://img/x.png"},
            ],
        }
    ]
    normalized = client._normalize_to_openai_format(messages)
    assert normalized[0]["content"][0] == {"type": "text", "text": "问题"}
    assert normalized[0]["content"][1] == {"type": "image_url", "image_url": {"url": "https://img/x.png"}}


def test_image_service_signed_url_path_and_verify():
    service = ImageService()
    url = service.generate_signed_url(file_id="f1", image_id="page_0001", ttl=60)
    parsed = urlparse(url)

    assert "/api/knowledge/images/f1/page_0001.png" in parsed.path
    query = parse_qs(parsed.query)
    assert "expires" in query and "sig" in query

    expires = int(query["expires"][0])
    sig = query["sig"][0]
    assert service.verify_signature("f1", "page_0001", expires, sig)


def test_image_service_signed_url_uses_public_base_url():
    settings = get_settings()
    original_base = settings.app.public_api_base_url
    settings.app.public_api_base_url = "https://public.example.com"
    try:
        service = ImageService()
        url = service.generate_signed_url(file_id="f1", image_id="page_0001", ttl=60)
        assert url.startswith("https://public.example.com/api/knowledge/images/")
    finally:
        settings.app.public_api_base_url = original_base


def test_document_ingestor_pdf_fixed_chunking_uses_config():
    settings = get_settings()
    original_size = settings.rag.pdf_fixed_chunk_size
    original_overlap = settings.rag.pdf_fixed_chunk_overlap
    settings.rag.pdf_fixed_chunk_size = 10
    settings.rag.pdf_fixed_chunk_overlap = 2
    try:
        ingestor = DocumentIngestor(chroma_client=MagicMock())
        text = "[PAGE:1]\nabcdefghijklmnopqrstuvwxyz"
        chunks = ingestor._chunk_pdf_fixed(
            text=text,
            file_id="file_x",
            source_file="demo.pdf",
            workspace_id="ws_1",
        )
        assert len(chunks) >= 3
        assert chunks[0].metadata["page_numbers"] == [1]
        assert chunks[0].metadata["chunk_index"] == 0
        assert chunks[0].next_id == chunks[1].chunk_id
    finally:
        settings.rag.pdf_fixed_chunk_size = original_size
        settings.rag.pdf_fixed_chunk_overlap = original_overlap


def test_document_ingestor_pdf_fixed_chunking_empty_pages_returns_no_chunks():
    ingestor = DocumentIngestor(chroma_client=MagicMock())
    text = "[PAGE:1]\n\n[PAGE:2]\n   \n"
    chunks = ingestor._chunk_pdf_fixed(
        text=text,
        file_id="file_empty",
        source_file="empty.pdf",
        workspace_id="ws_1",
    )
    assert chunks == []


def test_async_embedding_embed_texts_rejects_empty_input():
    client = AsyncEmbeddingClient.__new__(AsyncEmbeddingClient)
    client.model = "text-embedding-v3"
    client.dimensions = 1024

    try:
        asyncio.run(client.embed_texts([]))
        assert False, "expected ValueError for empty texts"
    except ValueError as e:
        assert "非空文本列表" in str(e)


def test_build_citation_file_slot_supports_page_query():
    settings = get_settings()
    original_enable = settings.rag.citation_enable_page_link
    original_default_page = settings.rag.citation_default_page
    original_page_key = settings.rag.citation_page_param_key

    settings.rag.citation_enable_page_link = True
    settings.rag.citation_default_page = 1
    settings.rag.citation_page_param_key = "page"
    try:
        assert _build_citation_file_slot("file_1", "pdf", 12) == "file_1?page=12"
        assert _build_citation_file_slot("file_1", "pdf", None) == "file_1?page=1"
        assert _build_citation_file_slot("file_1", "docx", 5) == "file_1?page=5"
        assert _build_citation_file_slot("file_1", "docx", None) == "file_1"
        assert _build_citation_file_slot("file_1", "doc", 9) == "file_1?page=9"
        assert _build_citation_file_slot("file_1", "doc", None) == "file_1"
        settings.rag.citation_enable_page_link = False
        assert _build_citation_file_slot("file_1", "pdf", 12) == "file_1"
    finally:
        settings.rag.citation_enable_page_link = original_enable
        settings.rag.citation_default_page = original_default_page
        settings.rag.citation_page_param_key = original_page_key


def test_parse_page_numbers_handles_comma_string_inside_list():
    pages = _parse_page_numbers(["12, 13", "15"])
    assert pages == [12, 13, 15]


def test_synthesizer_collect_citations_keeps_same_file_different_pages():
    results = [
        {
            "result": (
                "A [[CITATION:file_1?page=1:report.pdf]] "
                "B [[CITATION:file_1?page=2:report.pdf]] "
                "C [[CITATION:file_1?page=1:report.pdf]]"
            )
        }
    ]

    citations = synthesizer_module._collect_citations(results, page_param_key="page")
    keys = {(item["clean_file_id"], item["page_number"], item["file_name"]) for item in citations}
    assert len(citations) == 2
    assert keys == {("file_1", 1, "report.pdf"), ("file_1", 2, "report.pdf")}


def test_direct_execute_rag_only_passes_meta(monkeypatch):
    async def _noop(*args, **kwargs):
        return None

    async def _fake_execute_worker(**kwargs):
        return {
            "output": "ok",
            "memory_update": {},
            "meta": {
                "mm_evidence": {
                    "images": [{"file_id": "f1", "page_number": 3, "url": "https://img/3"}]
                }
            },
        }

    monkeypatch.setattr(direct_execute_module, "emit_step_update", _noop)
    monkeypatch.setattr(direct_execute_module, "emit_plan_complete", _noop)
    monkeypatch.setattr(direct_execute_module, "_execute_worker", _fake_execute_worker)

    output = asyncio.run(
        direct_execute_module.direct_execute_node(
            {
                "execution_mode": "rag_only",
                "user_query": "测试",
                "user_context": {},
                "session_id": "s1",
                "memory_dfs": {},
                "round_index": 0,
                "messages": [],
            }
        )
    )

    result = output["execution_results"][0]
    assert result["worker"] == "doc_worker"
    assert result["meta"]["mm_evidence"]["images"][0]["page_number"] == 3


def test_document_ingestor_persist_document_images_records(monkeypatch):
    class _FakeResult:
        def scalar_one_or_none(self):
            return None

    class _FakeSession:
        def __init__(self):
            self.executed = []
            self.added = []

        async def execute(self, stmt):
            self.executed.append(stmt)
            return _FakeResult()

        def add_all(self, rows):
            self.added.extend(rows)

    class _FakeManager:
        def __init__(self, session):
            self._session = session

        @asynccontextmanager
        async def session_scope(self):
            yield self._session

    class _Img:
        def __init__(self):
            self.id = "001"
            self.local_path = "/tmp/001.png"
            self.page_number = 5
            self.bbox = {"x1": 1, "y1": 2, "x2": 3, "y2": 4}
            self.phash = "abc"
            self.width = 800
            self.height = 600

    fake_session = _FakeSession()
    monkeypatch.setattr(
        document_ingestor_module,
        "get_async_db_manager",
        lambda: _FakeManager(fake_session),
    )

    ingestor = DocumentIngestor(chroma_client=MagicMock())
    asyncio.run(ingestor._persist_document_images("file_x", [_Img()]))

    assert len(fake_session.executed) == 1
    assert len(fake_session.added) == 1
    saved = fake_session.added[0]
    assert saved.file_id == "file_x"
    assert saved.image_id == "001"
    assert saved.page_number == 5


def test_image_meta_api_returns_page_number():
    class _FakeResult:
        def __init__(self, row):
            self._row = row

        def first(self):
            return self._row

    class _FakeDb:
        async def execute(self, _stmt):
            return _FakeResult(("row-id", 9))

    user = User(
        id="u1",
        username="tester",
        role="admin",
        workspace_id="ws_1",
        permissions=["*"],
    )
    resp = asyncio.run(
        get_image_meta(
            file_id="f1",
            image_id="001",
            current_user=user,
            db=_FakeDb(),
        )
    )
    assert resp.page_number == 9


def test_image_meta_api_null_page_number_falls_back_to_one():
    class _FakeResult:
        def __init__(self, row):
            self._row = row

        def first(self):
            return self._row

    class _FakeDb:
        async def execute(self, _stmt):
            return _FakeResult(("row-id", None))

    user = User(
        id="u1",
        username="tester",
        role="admin",
        workspace_id="ws_1",
        permissions=["*"],
    )
    resp = asyncio.run(
        get_image_meta(
            file_id="f1",
            image_id="001",
            current_user=user,
            db=_FakeDb(),
        )
    )
    assert resp.page_number == 1


def test_image_meta_api_fallback_page_cache_id():
    class _FakeResult:
        def __init__(self, row=None, scalar=None):
            self._row = row
            self._scalar = scalar

        def first(self):
            return self._row

        def scalar_one_or_none(self):
            return self._scalar

    class _FakeDb:
        def __init__(self):
            self.calls = 0

        async def execute(self, _stmt):
            self.calls += 1
            if self.calls == 1:
                return _FakeResult(row=None)  # document_images 未命中
            return _FakeResult(scalar="f1")  # 文件存在且属于当前租户

    user = User(
        id="u1",
        username="tester",
        role="admin",
        workspace_id="ws_1",
        permissions=["*"],
    )
    resp = asyncio.run(
        get_image_meta(
            file_id="f1",
            image_id="page_0012",
            current_user=user,
            db=_FakeDb(),
        )
    )
    assert resp.page_number == 12


def test_image_meta_api_not_found_raises_404():
    class _FakeResult:
        def __init__(self, row=None, scalar=None):
            self._row = row
            self._scalar = scalar

        def first(self):
            return self._row

        def scalar_one_or_none(self):
            return self._scalar

    class _FakeDb:
        def __init__(self):
            self.calls = 0

        async def execute(self, _stmt):
            self.calls += 1
            return _FakeResult(row=None, scalar=None)

    user = User(
        id="u1",
        username="tester",
        role="admin",
        workspace_id="ws_1",
        permissions=["*"],
    )
    with pytest.raises(HTTPException) as exc_info:
        asyncio.run(
            get_image_meta(
                file_id="f1",
                image_id="page_0012",
                current_user=user,
                db=_FakeDb(),
            )
        )
    assert exc_info.value.status_code == 404


def test_image_id_normalization_supports_multiple_suffixes():
    assert knowledge_images_module._normalize_image_id("001.png") == "001"
    assert knowledge_images_module._normalize_image_id("001.jpg") == "001"
    assert knowledge_images_module._normalize_image_id("001.jpeg") == "001"
    assert knowledge_images_module._normalize_image_id("page_0002") == "page_0002"
