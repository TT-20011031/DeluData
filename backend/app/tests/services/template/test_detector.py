"""
模板空白检测器单元测试
"""
import pytest
import asyncio
import os
from unittest.mock import MagicMock, patch

# 导入被测模块
from app.services.template.detector import (
    BlankFieldDetector,
    get_blank_field_detector,
    CandidateField,
    FieldLocation,
)


class TestBlankFieldDetector:
    """空白检测器测试"""

    def setup_method(self):
        """每个测试前重置单例"""
        global _detector
        from app.services.template import detector
        detector._detector = None

    def test_singleton_pattern(self):
        """测试单例模式"""
        d1 = get_blank_field_detector()
        d2 = get_blank_field_detector()
        assert d1 is d2

    def test_candidate_field_model(self):
        """测试 CandidateField 模型"""
        field = CandidateField(
            key="f_0001",
            label="项目名称",
            type="text",
            location=FieldLocation(
                type="paragraph",
                paragraph_index=5,
                selected_text="____"
            ),
            confidence=0.85,
            source="underline_chars",
            context="项目名称："
        )
        assert field.key == "f_0001"
        assert field.confidence == 0.85
        assert field.location.type == "paragraph"

    def test_field_location_paragraph(self):
        """测试段落位置模型"""
        loc = FieldLocation(
            type="paragraph",
            paragraph_index=10,
            selected_text="________"
        )
        assert loc.type == "paragraph"
        assert loc.paragraph_index == 10
        assert loc.table_index is None

    def test_field_location_table_cell(self):
        """测试表格单元格位置模型"""
        loc = FieldLocation(
            type="table_cell",
            table_index=0,
            row_index=3,
            cell_index=1,
            selected_text=""
        )
        assert loc.type == "table_cell"
        assert loc.table_index == 0
        assert loc.row_index == 3

    def test_detector_initialization(self):
        """测试检测器初始化"""
        detector = BlankFieldDetector()
        assert detector._settings is None  # 懒加载
        assert detector._field_counter == 0

    def test_generate_key(self):
        """测试 key 生成"""
        detector = BlankFieldDetector()
        key1 = detector._generate_key()
        key2 = detector._generate_key()
        assert key1 == "f_0001"
        assert key2 == "f_0002"

    def test_clean_label_with_colon(self):
        """测试标签清理 - 带冒号"""
        detector = BlankFieldDetector()
        assert detector._clean_label("项目名称：") == "项目名称"
        assert detector._clean_label("姓名:") == "姓名"

    def test_clean_label_with_punctuation(self):
        """测试标签清理 - 带标点"""
        detector = BlankFieldDetector()
        assert detector._clean_label("联系电话、") == "联系电话"
        assert detector._clean_label("地址，") == "地址"

    def test_clean_label_too_long(self):
        """测试标签清理 - 过长文本"""
        detector = BlankFieldDetector()
        long_text = "这是一个非常非常长的文本内容超过二十个字符应该被忽略"
        assert detector._clean_label(long_text) == ""

    def test_extract_context(self):
        """测试上下文提取"""
        detector = BlankFieldDetector()
        text = "项目名称：________联系电话"
        context = detector._extract_context(text, 5, 13)  # ________ 的位置
        assert "项目名称" in context

    def test_build_context(self):
        """测试上下文构建"""
        detector = BlankFieldDetector()
        ctx = detector._build_context("前缀文字", "后缀文字")
        assert ctx == "前缀文字 | 后缀文字"
        
        ctx_no_suffix = detector._build_context("只有前缀", "")
        assert ctx_no_suffix == "只有前缀"


class TestDetectorIntegration:
    """集成测试（需要实际文件）"""

    @pytest.fixture
    def sample_docx_path(self, tmp_path):
        """创建测试用 docx 文件"""
        from docx import Document
        
        doc = Document()
        
        # 添加带下划线的段落
        doc.add_paragraph("姓名：________")
        doc.add_paragraph("联系电话：____________")
        doc.add_paragraph("（请填写）")
        
        # 添加表格
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "项目名称"
        table.rows[0].cells[1].text = ""  # 空单元格
        table.rows[1].cells[0].text = "备注"
        table.rows[1].cells[1].text = ""  # 空单元格
        
        filepath = tmp_path / "test_template.docx"
        doc.save(str(filepath))
        return str(filepath)

    @pytest.mark.asyncio
    async def test_detect_async(self, sample_docx_path):
        """测试异步检测"""
        detector = BlankFieldDetector()
        candidates = await detector.detect_async(sample_docx_path)
        
        assert isinstance(candidates, list)
        # 应该检测到多个候选
        assert len(candidates) > 0
        
        # 验证字段结构
        for c in candidates:
            assert isinstance(c, CandidateField)
            assert c.key.startswith("f_")
            assert 0 <= c.confidence <= 1
            assert c.source in [
                "underline_chars", "bracket_blank", "run_underline",
                "paragraph_border", "empty_cell", "empty_cell_with_label"
            ]

    @pytest.mark.asyncio
    async def test_detect_underline_chars(self, sample_docx_path):
        """测试下划线检测"""
        detector = BlankFieldDetector()
        candidates = await detector.detect_async(sample_docx_path)
        
        underline_candidates = [c for c in candidates if c.source == "underline_chars"]
        assert len(underline_candidates) >= 2  # 至少应该检测到两个下划线

    @pytest.mark.asyncio
    async def test_detect_empty_table_cells(self, sample_docx_path):
        """测试空表格单元格检测"""
        detector = BlankFieldDetector()
        candidates = await detector.detect_async(sample_docx_path)
        
        table_candidates = [c for c in candidates if c.location.type == "table_cell"]
        assert len(table_candidates) >= 2  # 至少两个空单元格

    @pytest.mark.asyncio
    async def test_context_extraction(self, sample_docx_path):
        """测试上下文提取"""
        detector = BlankFieldDetector()
        candidates = await detector.detect_async(sample_docx_path)
        
        # 找到有上下文的候选
        with_context = [c for c in candidates if c.context]
        # 应该有一些候选有上下文
        assert len(with_context) > 0


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
