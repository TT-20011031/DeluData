import json
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

sys.path.insert(0, ".")

from app.agents.office.handlers.base import TaskContext
from app.agents.office.handlers.excel_creation import ExcelCreationHandler
from app.agents.office_worker import OfficeWorker
from app.api.chat.schemas import StartSessionRequest
import app.agents.office.handlers.excel_creation as excel_creation_module
import app.api.events as events_module
import app.services.skill_retriever as skill_retriever_module
import app.services.temp_artifact_service as temp_artifact_service_module


def test_start_session_request_accepts_skill_id_and_forbids_unknown_fields():
    payload = StartSessionRequest.model_validate(
        {
            "message": "执行",
            "session_id": "session-1",
            "skill_id": "skill-1",
        }
    )

    assert payload.skill_id == "skill-1"

    with pytest.raises(ValidationError):
        StartSessionRequest.model_validate({"message": "执行", "unexpected": True})


@pytest.mark.asyncio
async def test_retrieve_skill_for_query_raises_when_explicit_skill_is_missing(monkeypatch):
    class FakeSkillService:
        async def get_skill(self, _skill_id):
            return None

    monkeypatch.setattr(skill_retriever_module, "get_skill_service", lambda: FakeSkillService())

    with pytest.raises(ValueError, match="不存在或已删除"):
        await skill_retriever_module.retrieve_skill_for_query(
            query="执行",
            user_context=SimpleNamespace(workspace_id="ws-1"),
            skill_id="skill-missing",
        )


@pytest.mark.asyncio
async def test_retrieve_skill_for_query_raises_when_explicit_skill_crosses_workspace(monkeypatch):
    class FakeSkillService:
        async def get_skill(self, _skill_id):
            return SimpleNamespace(
                id="skill-1",
                title="跨工作区手册",
                workspace_id="other-ws",
                steps=[],
            )

    monkeypatch.setattr(skill_retriever_module, "get_skill_service", lambda: FakeSkillService())

    with pytest.raises(ValueError, match="当前工作区"):
        await skill_retriever_module.retrieve_skill_for_query(
            query="执行",
            user_context=SimpleNamespace(workspace_id="ws-1"),
            skill_id="skill-1",
        )


def test_office_worker_detects_excel_creation_from_query_and_filename():
    worker = OfficeWorker()

    excel_ctx = TaskContext(
        task_description="请生成 Excel 销售报表",
        sandbox_path="D:/tmp",
        session_id="s-1",
        output_filename=None,
    )
    assert worker._detect_task_mode(excel_ctx) == "excel_creation"

    excel_file_ctx = TaskContext(
        task_description="请生成本周汇总",
        sandbox_path="D:/tmp",
        session_id="s-2",
        output_filename="周报.xlsx",
    )
    assert worker._detect_task_mode(excel_file_ctx) == "excel_creation"


@pytest.mark.asyncio
async def test_excel_creation_handler_renders_high_quality_workbook(monkeypatch, tmp_path):
    class FakeLLM:
        async def chat(self, *_args, **_kwargs):
            return json.dumps(
                {
                    "filename": "销售报表.xlsx",
                    "workbook_title": "区域销售分析",
                    "sheets": [
                        {
                            "name": "销售报表",
                            "description": "按城市汇总的核心指标",
                            "freeze_header": True,
                            "columns": [
                                {"header": "城市", "key": "city", "type": "text"},
                                {"header": "销售额", "key": "sales", "type": "currency"},
                                {"header": "同比增长", "key": "growth", "type": "percentage"},
                            ],
                            "rows": [
                                {"city": "北京", "sales": 100000, "growth": 0.12},
                                {"city": "上海", "sales": 150000, "growth": 0.08},
                            ],
                            "summary": {
                                "label": "合计",
                                "formulas": {
                                    "sales": "sum",
                                    "growth": "average",
                                },
                            },
                            "chart": {
                                "type": "column",
                                "title": "城市销售额对比",
                                "category_key": "city",
                                "series": [
                                    {"name": "销售额", "value_key": "sales"},
                                ],
                            },
                        }
                    ],
                },
                ensure_ascii=False,
            )

        def _parse_json_from_text(self, text: str):
            return json.loads(text)

    class FakeArtifactService:
        async def save_doc(self, **kwargs):
            self.last_call = kwargs
            return kwargs

    async def fake_relevance(**_kwargs):
        return {"is_relevant": True, "confidence": 0.91}

    async def fake_emit_file_result(**_kwargs):
        return None

    fake_artifact_service = FakeArtifactService()
    monkeypatch.setattr(excel_creation_module, "get_async_llm", lambda: FakeLLM())
    monkeypatch.setattr(excel_creation_module, "evaluate_generation_data_relevance", fake_relevance)
    monkeypatch.setattr(
        ExcelCreationHandler,
        "_build_creation_context",
        lambda self, _ctx: "\n## mock data\n- 北京 100000\n- 上海 150000",
    )
    monkeypatch.setattr(
        temp_artifact_service_module,
        "get_temp_artifact_service",
        lambda: fake_artifact_service,
    )
    monkeypatch.setattr(events_module, "emit_file_result", fake_emit_file_result)

    handler = ExcelCreationHandler()
    ctx = TaskContext(
        task_description="生成 Excel 销售报表",
        sandbox_path=str(tmp_path),
        session_id="session-1",
        user_id="user-1",
        parent_step_id="step-1",
        user_context={"workspace_id": "ws-1", "user_id": "user-1"},
        memory_dfs={"dataset": [{"city": "北京", "sales": 100000}]},
        execution_results=[],
        round_index=2,
    )

    result = await handler.handle(ctx)

    assert result["success"] is True
    output_file = Path(result["output_files"][0]["path"])
    assert output_file.exists()
    assert output_file.suffix == ".xlsx"
    assert fake_artifact_service.last_call["file_kind"] == "excel"

    workbook = load_workbook(output_file)
    sheet = workbook["销售报表"]

    assert sheet.freeze_panes == "A4"
    assert len(sheet.tables) == 1
    assert sheet["A6"].value == "合计"
    assert "SUBTOTAL" in str(sheet["B6"].value).upper()
    assert "SUBTOTAL" in str(sheet["C6"].value).upper()
    assert len(sheet._charts) == 1
