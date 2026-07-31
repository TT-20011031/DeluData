"""
Word 图片提取顺序测试脚本

验证 python-docx 在混合排版（嵌入 vs 浮动）下的图片提取顺序是否与视觉顺序一致。

使用方法：
    python test_docx_image_order.py <测试文档路径.docx>
    
测试要点：
1. 嵌入式图片 (Inline) 的提取顺序
2. 浮动图片 (Floating) 的提取顺序
3. 混合排版场景下的顺序一致性
"""
import sys
from pathlib import Path
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from io import BytesIO
from PIL import Image
import imagehash


def analyze_docx_images(docx_path: str):
    """分析 Word 文档中的图片及其提取顺序"""
    
    print(f"\n{'='*60}")
    print(f"📄 分析文档: {docx_path}")
    print(f"{'='*60}")
    
    doc = Document(docx_path)
    
    # ========== 方法 1: 通过 rels 遍历（当前 ImageService 使用的方法）==========
    print(f"\n🔍 方法 1: 通过 document.part.rels 遍历")
    print("-" * 40)
    
    rels_images = []
    for rel_id, rel in doc.part.rels.items():
        if rel.reltype == RT.IMAGE:
            try:
                image_data = rel.target_part.blob
                img = Image.open(BytesIO(image_data))
                phash = str(imagehash.phash(img))[:8]
                rels_images.append({
                    "rel_id": rel_id,
                    "size": f"{img.width}x{img.height}",
                    "phash": phash,
                    "target": rel.target_ref
                })
            except Exception as e:
                print(f"  ⚠️ 解析失败: {rel_id} - {e}")
    
    for i, img_info in enumerate(rels_images, 1):
        print(f"  {i:03d}. {img_info['rel_id']:10} | {img_info['size']:10} | hash={img_info['phash']}")
    
    # ========== 方法 2: 遍历段落中的 inline shapes ==========
    print(f"\n🔍 方法 2: 遍历段落中的 InlineShapes (嵌入式图片)")
    print("-" * 40)
    
    inline_count = 0
    for para_idx, para in enumerate(doc.paragraphs):
        # 检查段落中的 run
        for run in para.runs:
            # 检查 inline shapes
            inline_shapes = run._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline')
            for shape in inline_shapes:
                inline_count += 1
                # 尝试获取图片引用
                blip = shape.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
                if blip is not None:
                    embed_id = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    print(f"  {inline_count:03d}. 段落 {para_idx+1:3d} | rel_id={embed_id}")
    
    if inline_count == 0:
        print("  (未找到嵌入式图片)")
    
    # ========== 方法 3: 遍历浮动图片 (anchor) ==========
    print(f"\n🔍 方法 3: 遍历 Anchor 元素 (浮动图片)")
    print("-" * 40)
    
    float_count = 0
    for para_idx, para in enumerate(doc.paragraphs):
        for run in para.runs:
            anchors = run._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor')
            for anchor in anchors:
                float_count += 1
                blip = anchor.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
                if blip is not None:
                    embed_id = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    print(f"  {float_count:03d}. 段落 {para_idx+1:3d} | rel_id={embed_id} | ⚠️ 浮动")
    
    if float_count == 0:
        print("  (未找到浮动图片)")
    
    # ========== 总结 ==========
    print(f"\n📊 统计总结")
    print("-" * 40)
    print(f"  • rels 中的图片总数: {len(rels_images)}")
    print(f"  • 嵌入式图片 (Inline): {inline_count}")
    print(f"  • 浮动图片 (Anchor): {float_count}")
    
    if float_count > 0:
        print(f"\n  ⚠️ 警告: 检测到 {float_count} 张浮动图片!")
        print(f"     浮动图片的提取顺序可能与视觉顺序不一致。")
        print(f"     建议用户将图片设为 '嵌入型/与文字排列' 模式。")
    
    if len(rels_images) != (inline_count + float_count):
        print(f"\n  ⚠️ 警告: rels 图片数 ({len(rels_images)}) ≠ 实际引用数 ({inline_count + float_count})")
        print(f"     可能存在未使用的图片资源或重复引用。")
    
    # ========== 建议的排序策略 ==========
    print(f"\n💡 建议的排序策略")
    print("-" * 40)
    print("""
    当前 ImageService 使用 rels 遍历，顺序不确定。
    
    建议改为「段落顺序遍历」：
    1. 遍历所有段落
    2. 对每个段落，先收集 inline shapes，再收集 anchors
    3. 按照「段落顺序 + 段内位置」生成最终 ID
    
    这样可以保证：用户在第 3 段插入的图片，
    无论是嵌入还是浮动，都会被编号为较大的序号。
    """)


def main():
    if len(sys.argv) < 2:
        # 如果没有参数，创建测试文档并分析
        print("用法: python test_docx_image_order.py <文档路径.docx>")
        print("\n没有提供文档路径，将创建测试文档...")
        
        # 创建简易测试 Word 文档
        from docx.shared import Inches
        
        test_doc = Document()
        test_doc.add_heading("图片顺序测试文档", 0)
        
        test_doc.add_paragraph("[IMAGE:001]")
        test_doc.add_paragraph("这是第一张图片的描述...")
        
        test_doc.add_paragraph("[IMAGE:002]")
        test_doc.add_paragraph("这是第二张图片的描述...")
        
        test_path = "./test_image_order.docx"
        test_doc.save(test_path)
        print(f"已创建测试文档: {test_path}")
        print("请手动在此文档中插入图片后重新运行此脚本。")
        return
    
    docx_path = sys.argv[1]
    if not Path(docx_path).exists():
        print(f"❌ 文件不存在: {docx_path}")
        return
    
    analyze_docx_images(docx_path)


if __name__ == "__main__":
    main()
