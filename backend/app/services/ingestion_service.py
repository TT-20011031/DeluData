"""
文档解析服务层 (Ingestion Service)

职责：文档的 ETL (Extract, Transform, Load) 流程
- 文档格式转换（docx -> markdown）
- 向量化与索引
- 自动引用提取（Phase 3）

注意：Session 通过 __init__ 注入，确保事务一致性
"""
import asyncio
import logging
from contextlib import asynccontextmanager
from pathlib import Path
from threading import Event
from typing import Any, Callable, Optional

from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.services.filesystem_service import FilesystemService
from app.models.common.context import UserContext
from app.api.events import event_queue, EventEnvelope, EventChannel
from app.models.common.enums import DocumentStatus
from app.core.storage.service import get_storage_service
from app.core.db.database import get_async_db_manager
from app.models.knowledge.graph import DocumentImage

logger = logging.getLogger(__name__)


class IngestionService:
    """
    文档解析服务
    
    Layer 3 (编排层) - 协调多服务完成 ETL 流程
    """
    
    def __init__(self, db: Optional[AsyncSession] = None):
        """
        初始化服务
        
        Args:
            db: 可选数据库会话。为 None 时将按需创建短事务会话，
                避免 OCR/向量化长流程持有连接。
        """
        self.db = db
        self.filesystem = FilesystemService(db) if db is not None else None
        self._doc_skill = None
    
    @property
    def doc_skill(self):
        """懒加载 DocSkill（延迟导入以避免循环依赖）"""
        if self._doc_skill is None:
            from app.skills.doc_skill import DocSkill
            self._doc_skill = DocSkill()
        return self._doc_skill

    @asynccontextmanager
    async def _filesystem_scope(self):
        """
        获取 FilesystemService：
        - 若当前实例携带会话，则复用（兼容依赖注入场景）
        - 若无会话，则按需创建短事务会话（后台任务稳态场景）
        """
        if self.filesystem is not None:
            yield self.filesystem
            return

        db_manager = get_async_db_manager()
        async with db_manager.session_scope() as session:
            yield FilesystemService(session)

    # ================= 文档转换 =================

    async def docx_to_markdown(self, docx_path: str) -> str:
        """
        将 Word 文件转换为 Markdown 格式
        
        使用线程池执行 CPU 密集型操作，避免阻塞 Event Loop
        """
        def _sync_parse():
            from docx import Document
            
            doc = Document(docx_path)
            lines = []
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    lines.append("")
                    continue
                
                style_name = para.style.name.lower() if para.style else ""
                
                if 'heading 1' in style_name:
                    lines.append(f"# {text}")
                elif 'heading 2' in style_name:
                    lines.append(f"## {text}")
                elif 'heading 3' in style_name:
                    lines.append(f"### {text}")
                elif 'list bullet' in style_name:
                    lines.append(f"- {text}")
                elif 'list number' in style_name:
                    lines.append(f"1. {text}")
                else:
                    lines.append(text)
            
            return '\n'.join(lines)
        
        return await asyncio.to_thread(_sync_parse)

    async def save_markdown_to_docx(self, docx_path: str, markdown_content: str):
        """
        [v2.2] 异步将 Markdown 内容保存回 Word 文件
        
        使用 asyncio.to_thread 避免阻塞事件循环
        """
        await asyncio.to_thread(self._save_markdown_to_docx_sync, docx_path, markdown_content)
    
    def _save_markdown_to_docx_sync(self, docx_path: str, markdown_content: str):
        """
        [同步] 将 Markdown 内容保存回 Word 文件
        
        注意：简单解析，会丢失复杂样式
        """
        from docx import Document
        
        doc = Document()
        
        lines = markdown_content.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                doc.add_paragraph()
            elif line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line[0].isdigit() and '. ' in line[:4]:
                doc.add_paragraph(line.split('. ', 1)[1], style='List Number')
            else:
                doc.add_paragraph(line)
        
        doc.save(docx_path)
        logger.info(f"Markdown 转 Word 完成: {docx_path}")

    # ================= ETL 流程 =================

    async def process_document(
        self,
        doc_id: str,
        file_path: str,
        user_context: UserContext,
        filename: str,
        target_dept_id: Optional[str] = None,
        visibility: str = "dept",
        cancel_event: Optional[Event] = None,
        progress_callback: Optional[Callable[[str, int, Optional[dict[str, Any]]], Any]] = None,
    ):
        """
        文档处理主流程（向量化 + 索引）
        
        Args:
            doc_id: 文档 ID
            file_path: 文件路径
            user_context: 用户上下文
            filename: 文件名
            target_dept_id: 目标部门 ID
            visibility: 可见范围 (public | dept | private)
        """
        logger.info(f"开始处理文档: {doc_id}")

        if cancel_event is not None and cancel_event.is_set():
            logger.info("文档处理已取消（进入流程前）: %s", doc_id)
            return
        
        try:
            # 1. 向量化
            result = await self.doc_skill.ingest_document(
                file_path=file_path,
                user_context=user_context,
                file_id=doc_id,
                target_dept_id=target_dept_id,
                visibility=visibility,
                cancel_event=cancel_event,
                progress_callback=progress_callback,
            )
            
            # [关键修复] 检查入库结果，失败时标记 ERROR 并终止后续流程
            if not result.get('success', False):
                error_msg = result.get('error', '未知入库错误')
                if error_msg == "文档处理已取消":
                    logger.info("文档处理已取消: %s", doc_id)
                    return
                logger.warning(f"文档入库失败 [{doc_id}]: {error_msg}")
                async with self._filesystem_scope() as filesystem:
                    await filesystem.update_file_status(
                        doc_id, DocumentStatus.ERROR.value, error=error_msg
                    )
                await event_queue.broadcast(EventEnvelope(
                    channel=EventChannel.CONVERSATION,
                    type="message",
                    payload={
                        "type": "document_error",
                        "document_id": doc_id,
                        "message": f"文档处理失败: {error_msg}"
                    }
                ))
                return
            
            chunk_count = result.get('chunk_count', 0)
            
            # 2. 更新 MySQL 状态
            async with self._filesystem_scope() as filesystem:
                await filesystem.update_file_status(
                    doc_id,
                    DocumentStatus.INDEXED.value,
                    chunk_count=chunk_count
                )

            # 3. SSE 通知
            await event_queue.broadcast(EventEnvelope(
                channel=EventChannel.CONVERSATION,
                type="message",
                payload={
                    "type": "document_ready",
                    "document_id": doc_id,
                    "message": f"文档 '{filename}' 解析完成，共 {chunk_count} 个切片"
                }
            ))
            logger.info(f"文档处理完成: {doc_id}, chunks={chunk_count}")

            # 4. Wiki 自动编译：失败不影响主流程
            await self._maybe_enqueue_wiki_compile(
                file_id=doc_id,
                user_context=user_context,
                trigger_type="doc_upload",
            )

        except Exception as e:
            logger.error(f"文档处理失败 [{doc_id}]: {e}")
            async with self._filesystem_scope() as filesystem:
                await filesystem.update_file_status(
                    doc_id, DocumentStatus.ERROR.value, error=str(e)
                )
            
            await event_queue.broadcast(EventEnvelope(
                channel=EventChannel.CONVERSATION,
                type="message",
                payload={
                    "type": "document_error",
                    "document_id": doc_id,
                    "message": f"文档处理失败: {str(e)}"
                }
            ))

    async def update_document_vectors(
        self,
        file_id: str,
        file_path: str,
        user_context: UserContext,
        regenerate_description: bool = False,
        visibility: str = "dept",
        target_dept_id: Optional[str] = None,
        cancel_event: Optional[Event] = None,
        progress_callback: Optional[Callable[[str, int, Optional[dict[str, Any]]], Any]] = None,
    ):
        """
        更新文档向量（重新索引）
        
        Args:
            file_id: 文件 ID
            file_path: 文件路径
            user_context: 用户上下文
            regenerate_description: 是否重新生成描述
        """
        if cancel_event is not None and cancel_event.is_set():
            logger.info("文档向量更新已取消（进入流程前）: %s", file_id)
            return

        try:
            # 重新向量化
            result = await self.doc_skill.ingest_document(
                file_path=file_path,
                user_context=user_context,
                file_id=file_id,
                visibility=visibility,
                target_dept_id=target_dept_id,
                cancel_event=cancel_event,
                progress_callback=progress_callback,
            )
            if not result.get('success', False):
                error_msg = result.get('error', '未知入库错误')
                if error_msg == "文档处理已取消":
                    logger.info("文档向量更新已取消: %s", file_id)
                    return
                logger.warning(f"文档向量更新失败 [{file_id}]: {error_msg}")
                async with self._filesystem_scope() as filesystem:
                    await filesystem.update_file_status(
                        file_id, DocumentStatus.ERROR.value, error=error_msg
                    )
                return
            
            chunk_count = result.get('chunk_count', 0)
            
            # 可选：重新生成描述
            new_description = None
            if regenerate_description:
                new_description = await self._generate_file_description(file_path)
            
            async with self._filesystem_scope() as filesystem:
                await filesystem.update_file_status(
                    file_id,
                    DocumentStatus.INDEXED.value,
                    chunk_count=chunk_count,
                    description=new_description
                )
            logger.info(f"文档向量更新完成: {file_id}, chunks={chunk_count}")
            
            # SSE 通知
            await event_queue.broadcast(EventEnvelope(
                channel=EventChannel.CONVERSATION,
                type="message",
                payload={
                    "type": "document_ready",
                    "document_id": file_id,
                    "message": f"文档更新完成，共 {chunk_count} 个切片",
                    "new_description": new_description
                }
            ))

            # 文档重新入库后也触发 wiki 重编译（覆盖旧实体页内容）
            await self._maybe_enqueue_wiki_compile(
                file_id=file_id,
                user_context=user_context,
                trigger_type="doc_reindex",
            )

        except Exception as e:
            logger.error(f"文档向量更新失败 [{file_id}]: {e}")
            async with self._filesystem_scope() as filesystem:
                await filesystem.update_file_status(file_id, 'error', error=str(e))

    async def _generate_file_description(self, file_path: str) -> Optional[str]:
        """
        为文件生成描述（使用 LLM 摘要）
        """
        try:
            # [v2.3] 使用异步文件读取，避免阻塞事件循环
            from app.core.utils.file_utils import read_file_async
            
            storage_service = get_storage_service()
            suffix = Path(file_path).suffix
            async with storage_service.materialize(file_path, suffix=suffix) as local_file_path:
                content = await read_file_async(local_file_path, 3000)
            
            if not content.strip():
                return None
            
            # 使用 LLM 生成摘要
            from app.core.llm.async_llm import get_async_llm
            from app.config import get_settings
            from langchain_core.messages import HumanMessage
            
            settings = get_settings()
            llm = get_async_llm(model=settings.rag.summary_model)
            
            # [v2.3] 从配置获取 Prompt，解耦业务逻辑与配置
            prompt = settings.rag.summary_prompt.format(content=content)
            
            # [v2.3] 使用 HumanMessage 确保类型安全
            response = await llm.ainvoke([HumanMessage(content=prompt)])
            
            # [v2.3] 安全的返回值处理，兼容不同 LLM 类型
            result = response.content if hasattr(response, 'content') else str(response)
            return result[:200]
            
        except Exception as e:
            logger.error(f"生成文件描述失败: {e}")
            return None

    async def _maybe_enqueue_wiki_compile(
        self,
        *,
        file_id: str,
        user_context: UserContext,
        trigger_type: str = "doc_upload",
    ) -> None:
        """文档入库成功后入队 Wiki 编译任务。

        - 受 `wiki.enabled` + `wiki.auto_compile_on_ingest` 双开关控制
        - 入队失败不影响主流程，仅记日志
        - 入队幂等（同一 workspace + trigger_type 已 pending/running 时合并 file_ids）
        """
        try:
            from app.config import get_settings
            settings = get_settings()
            wiki_settings = settings.wiki
            if not (wiki_settings.enabled and wiki_settings.auto_compile_on_ingest):
                return

            from app.services.wiki_compile_queue_service import get_wiki_compile_queue

            queue = get_wiki_compile_queue()
            await queue.enqueue_task(
                workspace_id=user_context.workspace_id,
                trigger_type=trigger_type,
                user_id=user_context.user_id,
                payload={
                    "file_ids": [file_id],
                    "scope": "files",
                },
                deduplicate=True,
            )
            logger.info(
                "[Wiki] 入队编译任务: workspace=%s file=%s trigger=%s",
                user_context.workspace_id,
                file_id,
                trigger_type,
            )
        except Exception as exc:  # noqa: BLE001
            logger.warning(
                "[Wiki] 入队编译任务失败（不影响主流程）: workspace=%s file=%s err=%s",
                getattr(user_context, "workspace_id", None),
                file_id,
                exc,
            )

    async def pre_cleanup_document_data(
        self,
        *,
        file_id: str,
        user_context: UserContext,
    ) -> None:
        """
        任务重试/接管前的前置清理（一期：从头重跑）。

        约束：
        - 向量与关联数据清理放在单一 try/except 中统一失败处理
        - 任一步骤失败则抛出异常，由上层任务直接标记 failed/retry
        """
        try:
            # 1) 向量数据清理（Chroma + 检索缓存）
            vector_cleanup_ok = await self.doc_skill.delete_document(
                file_id,
                user_context,
                raise_on_error=True,
            )
            if not vector_cleanup_ok:
                logger.info(
                    "pre_cleanup: 未发现可清理的向量数据或无需清理 file_id=%s",
                    file_id,
                )

            # 2) 关系数据清理（document_images）
            db_manager = get_async_db_manager()
            async with db_manager.session_scope() as session:
                await session.execute(
                    delete(DocumentImage).where(
                        DocumentImage.file_id == file_id,
                    )
                )

            # 3) 中间产物清理（派生图片目录）
            from app.core.utils.image_service import get_image_service

            await asyncio.to_thread(get_image_service().delete_file_images, file_id)
        except Exception as e:
            logger.error("pre_cleanup 失败: file_id=%s err=%s", file_id, e, exc_info=True)
            raise

