"""
模板渲染器

职责：
- 模板渲染（docx/xlsx）
- 模板编译（插入占位符）

设计原则遵循：
- Async First: 使用 asyncio.to_thread 异步化阻塞操作
"""
import os
import re
import uuid
import asyncio
import logging
from typing import List, Dict, Any, Optional, Tuple

from app.models.config.template import Template

logger = logging.getLogger(__name__)


class TemplateRenderer:
    """模板渲染器"""

    async def render_to_temp(
        self,
        template: Template,
        context: Dict[str, Any]
    ) -> str:
        """
        渲染模板到临时文件

        Args:
            template: 模板实例
            context: 渲染上下文

        Returns:
            临时文件路径
        """
        import tempfile
        output_dir = tempfile.gettempdir()
        output_filename = f"test_{str(uuid.uuid4())[:8]}.{template.file_type}"
        output_path = os.path.join(output_dir, output_filename)

        await asyncio.to_thread(
            self._render_sync,
            template,
            context,
            output_path
        )
        return output_path

    async def render_to_path(
        self,
        template: Template,
        context: Dict[str, Any],
        output_path: str
    ) -> str:
        """
        渲染模板到指定路径
        
        [Async First] 使用 asyncio.to_thread 异步化阻塞操作

        Args:
            template: 模板实例
            context: 渲染上下文
            output_path: 输出文件路径

        Returns:
            输出路径
        """
        await asyncio.to_thread(
            self._render_sync,
            template,
            context,
            output_path
        )
        return output_path

    def _render_sync(
        self,
        template: Template,
        context: Dict[str, Any],
        output_path: str
    ) -> str:
        """
        同步渲染模板

        Args:
            template: 模板实例
            context: 渲染上下文
            output_path: 输出路径

        Returns:
            输出路径
        """
        file_path = template.file_path
        file_type = template.file_type

        if file_type == 'docx':
            self._render_docx(template, context, output_path)
        elif file_type == 'xlsx':
            self._render_xlsx(template, context, output_path)

        return output_path

    def _render_docx(
        self,
        template: Template,
        context: Dict[str, Any],
        output_path: str
    ):
        """渲染 Word 文档"""
        from docxtpl import DocxTemplate, InlineImage
        from docx.shared import Mm

        doc = DocxTemplate(template.file_path)
        render_context = {}
        schema = template.variables_schema or {}

        for k, v in context.items():
            var_def = schema.get(k, {})
            if var_def.get('type') == 'image' and v and os.path.exists(v):
                render_context[k] = InlineImage(doc, v, width=Mm(140))
            else:
                render_context[k] = v

        doc.render(render_context)
        doc.save(output_path)

    def _render_xlsx(
        self,
        template: Template,
        context: Dict[str, Any],
        output_path: str
    ):
        """渲染 Excel 表格"""
        import openpyxl
        from openpyxl.drawing.image import Image as XLImage

        wb = openpyxl.load_workbook(template.file_path)
        ws = wb.active
        schema = template.variables_schema or {}

        # 1. 文本/数字替换
        pattern = re.compile(r'\{\{(\w+)\}\}')
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    matches = pattern.findall(cell.value)
                    for match in matches:
                        if match in context:
                            val = context[match]
                            cell.value = cell.value.replace(f'{{{{{match}}}}}', str(val))

        # 2. 图片锚点处理
        for var_name, var_def in schema.items():
            if var_def.get('type') == 'image' and var_name in context:
                img_path = context[var_name]
                if img_path and os.path.exists(img_path):
                    placeholder = f'{{{{{var_name}}}}}'
                    for row in ws.iter_rows():
                        for cell in row:
                            if cell.value == placeholder:
                                img = XLImage(img_path)
                                ws.add_image(img, cell.coordinate)
                                cell.value = None
                                break

        # 3. 动态表格处理
        for var_name, var_def in schema.items():
            if var_def.get('type') == 'table' and var_name in context:
                table_data = context[var_name]
                if isinstance(table_data, list) and len(table_data) > 0:
                    start_row = None
                    for row_idx, row in enumerate(ws.iter_rows(), 1):
                        for cell in row:
                            if cell.value and f'{{{{{var_name}}}}}' in str(cell.value):
                                start_row = row_idx
                                cell.value = None
                                break
                        if start_row:
                            break

                    if start_row:
                        for i, row_data in enumerate(table_data):
                            for j, (key, val) in enumerate(row_data.items()):
                                ws.cell(row=start_row + i, column=j + 1, value=val)

        wb.save(output_path)

    # ========== 模板编译 ==========

    async def compile(
        self,
        template: Template,
        bindings: List[Dict[str, Any]],
        output_name: Optional[str] = None
    ) -> Tuple[str, str]:
        """
        编译模板（插入占位符）

        Args:
            template: 模板实例
            bindings: 绑定数据列表
            output_name: 输出文件名

        Returns:
            (输出路径, 输出文件名)
        """
        base_dir = os.path.dirname(template.file_path)
        base_name = output_name or f"compiled_{template.id}"
        compiled_name = base_name if base_name.endswith(".docx") else f"{base_name}.docx"
        output_path = os.path.join(base_dir, compiled_name)

        await asyncio.to_thread(
            self._compile_sync,
            template.file_path,
            bindings,
            output_path
        )
        return (output_path, compiled_name)

    def _compile_sync(
        self,
        file_path: str,
        bindings: List[Dict[str, Any]],
        output_path: str
    ) -> str:
        """
        同步编译模板

        Args:
            file_path: 源模板路径
            bindings: 绑定数据列表
            output_path: 输出路径

        Returns:
            输出路径
        """
        from docx import Document

        doc = Document(file_path)

        def _append_placeholder_to_paragraph(paragraph, placeholder_text: str) -> None:
            current_text = paragraph.text or ""
            if not current_text:
                paragraph.add_run(placeholder_text)
                return
            if placeholder_text in current_text:
                return
            trimmed = current_text.rstrip()
            sep = "" if trimmed.endswith((":", "：")) else " "
            suffix = f"{sep}{placeholder_text}"
            if paragraph.runs:
                paragraph.runs[-1].text += suffix
            else:
                paragraph.add_run(suffix)

        for b in bindings or []:
            key = b.get("key") or b.get("var") or b.get("name")
            if not key:
                continue
            placeholder = f"{{{{{key}}}}}"
            location = b.get("location") or {}
            ltype = location.get("type")
            selected_text = location.get("selected_text") or b.get("selected_text") or ""

            if ltype == "paragraph":
                pidx = location.get("paragraph_index")
                if isinstance(pidx, int) and 0 <= pidx < len(doc.paragraphs):
                    paragraph = doc.paragraphs[pidx]
                    if selected_text and self._replace_in_paragraph(paragraph, selected_text, placeholder):
                        continue
                    _append_placeholder_to_paragraph(paragraph, placeholder)

            elif ltype == "table_cell":
                ti = location.get("table_index")
                ri = location.get("row_index")
                ci = location.get("cell_index")
                if all(isinstance(x, int) for x in [ti, ri, ci]):
                    if 0 <= ti < len(doc.tables):
                        table = doc.tables[ti]
                        if 0 <= ri < len(table.rows):
                            row = table.rows[ri]
                            if 0 <= ci < len(row.cells):
                                cell = row.cells[ci]
                                if cell.paragraphs:
                                    paragraph = cell.paragraphs[0]
                                    if selected_text and self._replace_in_paragraph(paragraph, selected_text, placeholder):
                                        continue
                                    _append_placeholder_to_paragraph(paragraph, placeholder)
                                else:
                                    cell.text = (cell.text or "") + (placeholder if not cell.text else f" {placeholder}")

        doc.save(output_path)
        return output_path

    def _replace_in_paragraph(
        self,
        paragraph,
        target_text: str,
        replacement: str
    ) -> bool:
        """
        尝试在 run 内替换文本，保留样式

        Args:
            paragraph: python-docx 段落对象
            target_text: 目标文本
            replacement: 替换文本

        Returns:
            是否替换成功
        """
        if not target_text:
            return False
        for run in paragraph.runs:
            if target_text in run.text:
                run.text = run.text.replace(target_text, replacement)
                return True
        return False


# 单例
_renderer: TemplateRenderer = None


def get_template_renderer() -> TemplateRenderer:
    """获取渲染器单例"""
    global _renderer
    if _renderer is None:
        _renderer = TemplateRenderer()
    return _renderer
