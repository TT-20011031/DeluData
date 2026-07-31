"""
模板转换器

职责：
- docx 转 HTML（通过 Gotenberg）
- 注入 data-id 到 HTML 元素
- 提取文档结构映射

设计原则遵循：
- Async First: HTTP 调用使用 httpx 异步
"""
import os
import re
import logging
from typing import List, Dict, Any, Optional

from app.config import get_settings

logger = logging.getLogger(__name__)


class TemplateConverter:
    """模板转换器"""

    def __init__(self):
        self._gotenberg_settings = None

    @property
    def gotenberg_settings(self):
        """懒加载 Gotenberg 配置"""
        if self._gotenberg_settings is None:
            self._gotenberg_settings = get_settings().gotenberg
        return self._gotenberg_settings

    async def convert_docx_to_html(self, file_path: str) -> Optional[str]:
        """
        使用 Gotenberg 服务将 docx 转换为 HTML

        Args:
            file_path: docx 文件路径

        Returns:
            HTML 字符串，转换失败返回 None
        """
        gotenberg = self.gotenberg_settings
        if not gotenberg.url:
            return None

        base_url = gotenberg.url.rstrip("/")
        url = base_url if base_url.endswith("/convert") else f"{base_url}/convert"
        auth = (gotenberg.username, gotenberg.password) if gotenberg.username and gotenberg.password else None

        try:
            import httpx
            timeout = httpx.Timeout(gotenberg.timeout_seconds)
            async with httpx.AsyncClient(timeout=timeout) as client:
                with open(file_path, "rb") as f:
                    files = {
                        "file": (
                            os.path.basename(file_path),
                            f,
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                        )
                    }
                    response = await client.post(url, files=files, auth=auth)

            if response.status_code >= 400:
                logger.warning(f"HTML 预览失败: {response.status_code} {response.text[:200]}")
                return None

            try:
                payload = response.json()
            except Exception as e:
                logger.warning(f"HTML 预览解析失败: {e}")
                return None

            html_text = payload.get("html") if isinstance(payload, dict) else None
            if not html_text:
                logger.warning("HTML 预览返回空内容")
                return None
            if "<" not in html_text:
                logger.warning("HTML 预览返回内容疑似非HTML")
                return None

            return html_text
        except Exception as e:
            logger.warning(f"HTML 预览异常: {e}")
            return None

    def inject_data_ids_into_html(
        self,
        html: str,
        mapping: List[Dict[str, Any]]
    ) -> str:
        """
        将 data-id 注入到 HTML 元素中

        采用映射索引方式，确保 HTML 与编译索引一致

        Args:
            html: 原始 HTML
            mapping: 元素映射列表（来自 docx 解析）

        Returns:
            注入 data-id 后的 HTML
        """
        try:
            from bs4 import BeautifulSoup
        except Exception:
            logger.warning("模板预览缺少 bs4 依赖，无法注入 data-id")
            return html

        soup = BeautifulSoup(html, "html.parser")

        # ========== 按映射为所有段落注入 data-id ==========
        # 获取非表格内的段落元素
        paragraphs = [p for p in soup.find_all("p") if not p.find_parent("table")]
        
        # 如果没有 <p>，尝试其他块级元素
        if not paragraphs:
            paragraphs = [
                el for el in soup.find_all(["div", "section", "article"])
                if not el.find_parent("table")
            ]

        # 解析 mapping 中的段落 ID
        paragraph_ids = [
            item["id"] for item in mapping if item.get("type") == "paragraph"
        ]

        # 按映射顺序注入 data-id（不足则不注入）
        for idx, p in enumerate(paragraphs):
            if idx < len(paragraph_ids):
                p["data-id"] = paragraph_ids[idx]

        # ========== 按映射为所有表格单元格注入 data-id ==========
        cell_id_map: Dict[int, Dict[int, Dict[int, str]]] = {}
        for item in mapping:
            if item.get("type") != "table_cell":
                continue
            ti = item.get("table_index")
            ri = item.get("row_index")
            ci = item.get("cell_index")
            if not all(isinstance(x, int) for x in [ti, ri, ci]):
                continue
            cell_id_map.setdefault(ti, {}).setdefault(ri, {})[ci] = item["id"]

        tables = soup.find_all("table")
        
        for ti, table in enumerate(tables):
            rows = table.find_all("tr")
            for ri, row in enumerate(rows):
                cells = row.find_all(["td", "th"])
                for ci, cell in enumerate(cells):
                    mapped_id = cell_id_map.get(ti, {}).get(ri, {}).get(ci)
                    if mapped_id:
                        cell["data-id"] = mapped_id

        # 记录映射不匹配的情况
        if len(paragraphs) != len(paragraph_ids):
            logger.warning(
                f"HTML 段落数量与 docx 映射不一致: html={len(paragraphs)} docx={len(paragraph_ids)}"
            )
        html_cells = len([c for c in soup.find_all(["td", "th"])])
        mapped_cells = len([item for item in mapping if item.get("type") == "table_cell"])
        if html_cells != mapped_cells:
            logger.warning(
                f"HTML 单元格数量与 docx 映射不一致: html={html_cells} docx={mapped_cells}"
            )

        # 统计注入结果
        total_paragraphs = len([p for p in soup.find_all(True) if p.has_attr("data-id") and p.name == "p"])
        total_cells = len([c for c in soup.find_all(["td", "th"]) if c.has_attr("data-id")])
        
        logger.info(f"模板预览注入: 段落={total_paragraphs} 表格单元格={total_cells}")
        return str(soup)

    def _build_docx_mapping(self, file_path: str) -> List[Dict[str, Any]]:
        """
        使用 python-docx 生成映射（与编译索引一致）
        """
        try:
            from docx import Document
        except Exception:
            logger.warning("缺少 python-docx 依赖，无法生成 docx 映射")
            return []

        doc = Document(file_path)
        mapping: List[Dict[str, Any]] = []

        # 段落（仅文档主体，表格内段落在表格映射中）
        for idx, p in enumerate(doc.paragraphs):
            mapping.append({
                "id": f"p-{idx}",
                "type": "paragraph",
                "paragraph_index": idx,
                "text": p.text or ""
            })

        # 表格单元格
        for ti, table in enumerate(doc.tables):
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    mapping.append({
                        "id": f"t{ti}-r{ri}-c{ci}",
                        "type": "table_cell",
                        "table_index": ti,
                        "row_index": ri,
                        "cell_index": ci,
                        "text": cell.text or ""
                    })

        return mapping

    async def docx_to_html_with_mapping(self, file_path: str) -> Dict[str, Any]:
        """
        转换 docx 为 HTML 并返回元素映射

        Args:
            file_path: docx 文件路径

        Returns:
            {"html": str, "mapping": List[Dict]}

        Raises:
            RuntimeError: 转换失败
        """
        try:
            from bs4 import BeautifulSoup
        except Exception:
            raise RuntimeError("依赖 bs4 模块未安装")

        # 先转换 HTML
        html_text = await self.convert_docx_to_html(file_path)
        if not html_text:
            raise RuntimeError("HTML_CONVERT_UNAVAILABLE")

        # 生成与编译一致的 docx 映射
        mapping = self._build_docx_mapping(file_path)
        if not mapping:
            raise RuntimeError("DOCX_MAPPING_FAILED")

        # 注入 data-id，保证 HTML 与编译索引一致
        html_text = self.inject_data_ids_into_html(html_text, mapping)

        if 'data-id' not in html_text:
            raise RuntimeError("HTML_MAPPING_FAILED")

        logger.info(f"模板预览生成: mapping 数量={len(mapping)}")
        return {"html": html_text, "mapping": mapping}

    def to_dense_html(
        self, html: str, mapping: List[Dict[str, Any]]
    ) -> tuple[str, List[Dict[str, Any]]]:
        """
        将 Gotenberg 输出的 HTML 转换为干净的 Dense HTML
        
        功能：
        1. Flattening: 移除无语义容器（div/span），保留语义结构
        2. 视觉特征转译: 将下划线/占位符转为显式 <blank-hint> 标记
        3. Sub-ID 注入: 为每个 blank-hint 生成唯一子 ID
        4. 语义空格: 确保块级元素解包时不造成文本粘连
        
        Args:
            html: 原始 HTML（已含 data-id）
            mapping: 元素映射列表
            
        Returns:
            (dense_html, updated_mapping) - 清洗后的 HTML 和更新的映射
        """
        try:
            from bs4 import BeautifulSoup, NavigableString
        except ImportError:
            logger.warning("缺少 bs4 依赖，无法执行 Dense HTML 转换")
            return html, mapping
        
        soup = BeautifulSoup(html, "html.parser")
        blank_counter = 0
        blank_hints: List[Dict[str, Any]] = []
        
        # ========== Step 1: 视觉特征转译（在清洗前进行，避免丢失样式信息）==========
        blank_counter = self._convert_visual_blanks(soup, blank_hints, blank_counter)
        
        # ========== Step 2: 检测文本占位符（____、(  )、[请填写] 等）==========
        blank_counter = self._convert_placeholder_text(soup, blank_hints, blank_counter)
        
        # ========== Step 3: Flattening - 移除无语义容器 ==========
        self._flatten_semantic_containers(soup)
        
        # ========== Step 4: 移除所有 style/class 属性（已转译为语义标记）==========
        self._strip_style_attributes(soup)
        
        # 更新 mapping，添加 blank hints
        updated_mapping = mapping.copy()
        for hint in blank_hints:
            updated_mapping.append({
                "id": hint["id"],
                "type": "blank_hint",
                "parent_id": hint["parent_id"],
                "hint_type": hint["hint_type"],
                "text": "",
                "selected_text": hint.get("selected_text", "")
            })
        
        dense_html = str(soup)
        logger.info(f"Dense HTML 转换完成: 识别到 {len(blank_hints)} 个填空提示")
        return dense_html, updated_mapping

    def _convert_visual_blanks(
        self, soup, blank_hints: List[Dict], counter: int
    ) -> int:
        """
        将视觉上暗示填空的元素转换为 <blank-hint> 标记
        
        识别规则：
        - 带下划线样式 (text-decoration: underline) 的空/少内容元素
        - 带底边框 (border-bottom) 的空元素
        """
        # 查找带下划线或底边框的元素
        underline_patterns = [
            '[style*="underline"]',
            '[style*="border-bottom"]',
            'u',  # HTML underline 标签
        ]
        
        for selector in underline_patterns:
            for el in soup.select(selector):
                text = el.get_text(strip=True)
                # 空内容或仅含空格/下划线/占位符
                if not text or re.match(r'^[\s_\-—]+$', text):
                    # 找到最近的带 data-id 的父元素
                    parent_id = self._find_closest_data_id(el)
                    if parent_id:
                        blank_id = f"{parent_id}_blank_{counter}"
                        counter += 1
                        
                        # 创建 blank-hint 标签替换原元素
                        hint_tag = soup.new_tag("blank-hint")
                        hint_tag["id"] = blank_id
                        hint_tag["type"] = "underline"
                        hint_tag["parent-id"] = parent_id
                        el.replace_with(hint_tag)
                        
                        blank_hints.append({
                            "id": blank_id,
                            "parent_id": parent_id,
                            "hint_type": "underline",
                            "selected_text": text or ""
                        })
        
        return counter

    def _convert_placeholder_text(
        self, soup, blank_hints: List[Dict], counter: int
    ) -> int:
        """
        将文本占位符转换为 <blank-hint> 标记
        
        识别模式：
        - 连续下划线: ____, ———
        - 括号占位符: (    ), [   ], 【  】
        - 显式占位文本: [请填写], （待填）
        """
        placeholder_patterns = [
            (r'_{3,}', 'underline'),           # 3个及以上下划线
            (r'—{3,}', 'underline'),           # 3个及以上破折号
            (r'\(\s{2,}\)', 'parentheses'),    # 空括号 (  )
            (r'\[\s{2,}\]', 'brackets'),       # 空方括号 [  ]
            (r'【\s*】', 'brackets'),           # 空中括号 【】
            (r'\[请填写\]', 'placeholder'),    # 请填写
            (r'（待填）', 'placeholder'),       # 待填
            (r'\(待填\)', 'placeholder'),
        ]
        
        # 遍历所有文本节点
        from bs4 import NavigableString
        for text_node in soup.find_all(string=True):
            if not isinstance(text_node, NavigableString):
                continue
            
            original_text = str(text_node)
            parent = text_node.parent
            if not parent or parent.name in ['script', 'style', 'blank-hint']:
                continue
            
            parent_id = self._find_closest_data_id(parent)
            if not parent_id:
                continue
            
            # 检查是否匹配任何占位符模式
            modified = False
            new_content = []
            last_end = 0
            
            for pattern, hint_type in placeholder_patterns:
                for match in re.finditer(pattern, original_text):
                    modified = True
                    # 添加匹配前的文本
                    if match.start() > last_end:
                        new_content.append(original_text[last_end:match.start()])
                    
                    # 创建 blank-hint
                    blank_id = f"{parent_id}_blank_{counter}"
                    counter += 1
                    
                    hint_html = f'<blank-hint id="{blank_id}" type="{hint_type}" parent-id="{parent_id}"/>'
                    new_content.append(hint_html)
                    
                    blank_hints.append({
                        "id": blank_id,
                        "parent_id": parent_id,
                        "hint_type": hint_type,
                        "selected_text": match.group(0)
                    })
                    
                    last_end = match.end()
            
            if modified:
                # 添加剩余文本
                if last_end < len(original_text):
                    new_content.append(original_text[last_end:])
                
                # 用新内容替换原文本节点
                from bs4 import BeautifulSoup as BS
                new_soup = BS(''.join(new_content), 'html.parser')
                text_node.replace_with(new_soup)
        
        return counter

    def _flatten_semantic_containers(self, soup) -> None:
        """
        移除无语义的容器元素，保留内容
        
        注意：在 unwrap 时插入语义空格，避免文本粘连
        """
        # 要移除的无语义容器
        noise_tags = ['span', 'font']
        
        for tag_name in noise_tags:
            for tag in soup.find_all(tag_name):
                # 如果有 data-id 或 blank-hint，保留
                if tag.has_attr('data-id') or tag.name == 'blank-hint':
                    continue
                # unwrap 但保留内容
                tag.unwrap()
        
        # 对于块级容器（div），插入空格后再 unwrap
        for div in soup.find_all('div'):
            if div.has_attr('data-id'):
                continue
            # 在 div 后插入空格，防止文本粘连
            from bs4 import NavigableString
            div.insert_after(NavigableString(' '))
            div.unwrap()

    def _strip_style_attributes(self, soup) -> None:
        """移除所有 style 和 class 属性（视觉特征已转译为语义标记）"""
        for el in soup.find_all(True):
            # 保留 data-id, id, type, parent-id
            attrs_to_keep = {'data-id', 'id', 'type', 'parent-id', 
                           'colspan', 'rowspan', 'href', 'src'}
            attrs_to_remove = [
                attr for attr in el.attrs.keys() 
                if attr not in attrs_to_keep
            ]
            for attr in attrs_to_remove:
                del el[attr]

    def _find_closest_data_id(self, element) -> Optional[str]:
        """查找最近的带 data-id 的祖先元素"""
        current = element
        while current:
            if hasattr(current, 'has_attr') and current.has_attr('data-id'):
                return current['data-id']
            current = current.parent
        return None


# 单例
_converter: TemplateConverter = None


def get_template_converter() -> TemplateConverter:
    """获取转换器单例"""
    global _converter
    if _converter is None:
        _converter = TemplateConverter()
    return _converter
