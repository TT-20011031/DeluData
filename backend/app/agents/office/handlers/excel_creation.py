"""
Excel Creation Handler - 高质量 .xlsx 工作簿生成

使用 LLM 先规划结构化 workbook spec，再由 xlsxwriter 做确定性渲染。
"""
import copy
import json
import logging
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Literal, Optional

from pydantic import BaseModel, ConfigDict, Field

from app.config import get_settings
from app.core.llm.async_llm import get_async_llm
from app.core.llm.prompt_manager import get_prompt
from app.services.generation_data_guard import evaluate_generation_data_relevance
from app.services.generation_data_service import (
    build_query_focused_context,
    extract_query_focused_result_text,
)

from .base import TaskContext, TaskHandler

logger = logging.getLogger(__name__)


class ExcelColumnSpec(BaseModel):
    header: str
    key: str
    type: Literal["text", "integer", "number", "currency", "percentage", "date", "datetime"] = "text"


class ExcelChartSeriesSpec(BaseModel):
    name: Optional[str] = None
    value_key: str


class ExcelChartSpec(BaseModel):
    type: Literal["column", "bar", "line", "pie", "area", "scatter"] = "column"
    title: str = ""
    category_key: str
    series: list[ExcelChartSeriesSpec] = Field(default_factory=list)


class ExcelSummarySpec(BaseModel):
    label: str = "合计"
    formulas: dict[str, Literal["sum", "average", "count", "min", "max"]] = Field(default_factory=dict)


class ExcelSheetSpec(BaseModel):
    name: str
    description: str = ""
    freeze_header: bool = True
    columns: list[ExcelColumnSpec]
    rows: list[dict[str, Any]] = Field(default_factory=list)
    summary: Optional[ExcelSummarySpec] = None
    chart: Optional[ExcelChartSpec] = None


class ExcelWorkbookSpec(BaseModel):
    model_config = ConfigDict(extra="ignore")

    filename: Optional[str] = None
    workbook_title: str = ""
    sheets: list[ExcelSheetSpec]


class ExcelCreationHandler(TaskHandler):
    """
    Excel Creation 快速路径：LLM 生成 workbook spec -> xlsxwriter 确定性渲染
    """

    async def handle(self, ctx: TaskContext) -> Dict[str, Any]:
        await self.on_step(ctx, "生成Excel", "running", "正在规划工作簿结构...")

        relevance = await evaluate_generation_data_relevance(
            query=ctx.task_description,
            memory_dfs=ctx.memory_dfs,
            execution_results=ctx.execution_results,
            messages=ctx.messages,
        )
        if not relevance.get("is_relevant", True):
            reason = str(relevance.get("reason", "数据与问题不相关"))
            await self.on_step(ctx, "生成Excel", "invalid_data", "数据无效，已跳过生成")
            return {
                "success": True,
                "output": f"数据无效，已跳过 Excel 生成。原因：{reason}",
                "output_files": [],
                "quality_signal": {
                    "verdict": "pass",
                    "reason_code": "terminal_data_invalid",
                    "confidence": relevance.get("confidence", 0.9),
                    "retryable": False,
                },
                "meta": {
                    "worker_round": ctx.round_index,
                    "stop_reason": "terminal_data_invalid",
                    "step_status": "invalid_data",
                    "invalid_reason": reason,
                },
            }

        data_context = self._build_creation_context(ctx)
        llm = get_async_llm()
        settings = get_settings()
        system_prompt = get_prompt("excel_creation_system", module="office")
        user_prompt = get_prompt(
            "excel_creation_user",
            module="office",
            task_description=ctx.task_description,
            data_context=data_context,
        )

        try:
            raw = await llm.chat(
                [
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
                model=settings.llm.office_worker_model,
            )
            parsed = llm._parse_json_from_text(raw) if hasattr(llm, "_parse_json_from_text") else json.loads(raw)
            workbook_spec = self._coerce_workbook_spec(parsed)
        except Exception as exc:
            logger.error("ExcelCreationHandler: 解析 workbook spec 失败: %s", exc, exc_info=True)
            await self.on_step(ctx, "生成Excel", "error", "Excel 结构规划失败")
            return self._error_result(ctx, "Excel 工作簿结构解析失败", "excel_spec_invalid", retryable=True)

        if not self._has_meaningful_rows(workbook_spec):
            logger.warning("ExcelCreationHandler: 工作簿没有有效数据行")
            await self.on_step(ctx, "生成Excel", "error", "Excel 没有有效数据")
            return self._error_result(ctx, "Excel 工作簿未提取到有效数据，已阻止生成空白文件", "excel_empty_data", retryable=True)

        filename = self._resolve_filename(ctx, workbook_spec.filename)
        Path(ctx.sandbox_path).mkdir(parents=True, exist_ok=True)
        output_path = os.path.join(ctx.sandbox_path, filename)

        try:
            await self.on_step(ctx, "生成Excel", "running", "正在渲染高质量工作簿...")
            self._render_workbook(output_path, workbook_spec)
        except Exception as exc:
            logger.error("ExcelCreationHandler: 渲染失败: %s", exc, exc_info=True)
            await self.on_step(ctx, "生成Excel", "error", "Excel 渲染失败")
            return self._error_result(ctx, f"Excel 生成失败: {exc}", "excel_render_error", retryable=True)

        download_url = f"/api/files/download/{ctx.session_id}/{filename}"
        await self.on_step(ctx, "生成Excel", "done", "Excel 生成完成")

        try:
            from app.services.temp_artifact_service import get_temp_artifact_service

            user_ctx = ctx.user_context or {}
            workspace_id = str(user_ctx.get("workspace_id") or "default")
            user_id = str(user_ctx.get("user_id") or ctx.user_id or "anonymous")
            await get_temp_artifact_service().save_doc(
                workspace_id=workspace_id,
                user_id=user_id,
                session_id=ctx.session_id or "",
                title=filename[:50],
                download_url=download_url,
                file_name=filename,
                file_kind="excel",
            )
        except Exception as save_err:
            logger.warning("ExcelCreationHandler: 暂存箱保存失败(非致命): %s", save_err)

        if ctx.session_id:
            from app.api.events import emit_file_result

            await emit_file_result(
                session_id=ctx.session_id,
                step_id=ctx.parent_step_id,
                file_type="excel",
                file_name=filename,
                download_url=download_url,
                round_index=ctx.round_index,
            )

        return {
            "success": True,
            "output": f"已成功生成 Excel 工作簿 [{filename}]({download_url})。",
            "output_files": [{"name": filename, "path": output_path, "download_url": download_url}],
            "quality_signal": {
                "verdict": "pass",
                "reason_code": "excel_generated",
                "confidence": 0.92,
                "retryable": False,
            },
            "meta": {
                "worker_round": ctx.round_index,
                "stop_reason": "excel_generated",
            },
        }

    def _render_workbook(self, file_path: Any, workbook_spec: Any = None) -> None:
        if isinstance(file_path, (dict, ExcelWorkbookSpec)) and isinstance(workbook_spec, (str, os.PathLike)):
            file_path, workbook_spec = str(workbook_spec), file_path

        if not isinstance(file_path, (str, os.PathLike)):
            raise TypeError("file_path 必须是字符串或 PathLike")

        workbook_spec = self._coerce_workbook_spec(workbook_spec)
        try:
            import xlsxwriter
            from xlsxwriter.utility import xl_rowcol_to_cell
        except ModuleNotFoundError:
            logger.warning("ExcelCreationHandler: 未检测到 xlsxwriter，回退使用 openpyxl 渲染")
            self._render_workbook_with_openpyxl(str(file_path), workbook_spec)
            return

        workbook = xlsxwriter.Workbook(file_path)
        try:
            formats = self._build_formats(workbook)
            workbook_title = str(workbook_spec.workbook_title or "").strip()

            for index, sheet in enumerate(workbook_spec.sheets, start=1):
                sheet_name = self._sanitize_sheet_name(sheet.name or f"Sheet{index}", index)
                worksheet = workbook.add_worksheet(sheet_name)
                last_col = max(len(sheet.columns) - 1, 0)
                current_row = 0

                if workbook_title:
                    if last_col > 0:
                        worksheet.merge_range(current_row, 0, current_row, last_col, workbook_title, formats["title"])
                    else:
                        worksheet.write(current_row, 0, workbook_title, formats["title"])
                    current_row += 1

                description = str(sheet.description or "").strip()
                if description:
                    if last_col > 0:
                        worksheet.merge_range(current_row, 0, current_row, last_col, description, formats["description"])
                    else:
                        worksheet.write(current_row, 0, description, formats["description"])
                    current_row += 1

                data_matrix = [
                    [self._normalize_cell_value(row.get(column.key), column.type) for column in sheet.columns]
                    for row in sheet.rows
                ]

                header_row = current_row
                if data_matrix:
                    table_options = {
                        "style": "Table Style Medium 2",
                        "autofilter": True,
                        "data": data_matrix,
                        "columns": [
                            {
                                "header": column.header,
                                "format": formats[self._value_format_name(column.type)],
                            }
                            for column in sheet.columns
                        ],
                    }
                    table_last_row = header_row + len(data_matrix)
                    worksheet.add_table(header_row, 0, table_last_row, last_col, table_options)
                else:
                    table_last_row = header_row
                    worksheet.write_row(
                        header_row,
                        0,
                        [column.header for column in sheet.columns],
                        formats["header"],
                    )
                    worksheet.autofilter(header_row, 0, header_row, last_col)

                if sheet.freeze_header:
                    worksheet.freeze_panes(header_row + 1, 0)

                for col_idx, column in enumerate(sheet.columns):
                    width = self._calculate_column_width(column, sheet.rows)
                    worksheet.set_column(col_idx, col_idx, width, formats[self._value_format_name(column.type)])

                summary_row = table_last_row + 1
                if data_matrix and sheet.summary and sheet.summary.formulas:
                    worksheet.write(summary_row, 0, sheet.summary.label or "合计", formats["summary_label"])
                    for col_idx, column in enumerate(sheet.columns):
                        func = (sheet.summary.formulas or {}).get(column.key)
                        if not func:
                            continue
                        start_cell = xl_rowcol_to_cell(header_row + 1, col_idx)
                        end_cell = xl_rowcol_to_cell(table_last_row, col_idx)
                        formula = self._summary_formula(func, start_cell, end_cell)
                        worksheet.write_formula(
                            summary_row,
                            col_idx,
                            formula,
                            formats[self._summary_format_name(column.type)],
                        )
                    chart_anchor_row = summary_row + 2
                else:
                    chart_anchor_row = table_last_row + 2

                if data_matrix and sheet.chart and sheet.chart.series:
                    category_col_idx = self._find_column_index(sheet.columns, sheet.chart.category_key)
                    if category_col_idx is not None:
                        chart = workbook.add_chart({"type": self._chart_type(sheet.chart.type)})
                        categories = [
                            sheet_name,
                            header_row + 1,
                            category_col_idx,
                            table_last_row,
                            category_col_idx,
                        ]
                        for series_idx, series in enumerate(sheet.chart.series):
                            value_col_idx = self._find_column_index(sheet.columns, series.value_key)
                            if value_col_idx is None:
                                continue
                            chart.add_series(
                                {
                                    "name": series.name or sheet.columns[value_col_idx].header,
                                    "categories": categories,
                                    "values": [
                                        sheet_name,
                                        header_row + 1,
                                        value_col_idx,
                                        table_last_row,
                                        value_col_idx,
                                    ],
                                }
                            )
                            if sheet.chart.type == "pie":
                                break

                        if sheet.chart.title:
                            chart.set_title({"name": sheet.chart.title})
                        chart.set_legend({"position": "bottom"})
                        if sheet.chart.type != "pie":
                            chart.set_x_axis({"name": sheet.columns[category_col_idx].header})
                        worksheet.insert_chart(chart_anchor_row, last_col + 2, chart, {"x_scale": 1.25, "y_scale": 1.1})

                worksheet.conditional_format(
                    header_row,
                    0,
                    max(table_last_row, header_row),
                    last_col,
                    {
                        "type": "no_blanks",
                        "format": formats["table_body"],
                    },
                )

                worksheet.write(
                    max(chart_anchor_row, summary_row + 1 if data_matrix else header_row + 2),
                    0,
                    f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
                    formats["meta"],
                )
        finally:
            workbook.close()

    def _render_workbook_with_openpyxl(self, file_path: str, workbook_spec: ExcelWorkbookSpec) -> None:
        from openpyxl import Workbook
        from openpyxl.chart import Reference
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo

        workbook = Workbook()
        workbook.remove(workbook.active)

        thin = Side(style="thin", color="D0D7DE")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        title_font = Font(bold=True, size=16, color="17324D")
        description_font = Font(size=10, italic=True, color="5F6B7A")
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(fill_type="solid", fgColor="2F75B5")
        summary_fill = PatternFill(fill_type="solid", fgColor="D9E2F3")
        meta_font = Font(size=9, color="7A7A7A")
        centered = Alignment(horizontal="center", vertical="center")
        middle = Alignment(vertical="center")
        workbook_title = str(workbook_spec.workbook_title or "").strip()

        for index, sheet in enumerate(workbook_spec.sheets, start=1):
            sheet_name = self._sanitize_sheet_name(sheet.name or f"Sheet{index}", index)
            worksheet = workbook.create_sheet(sheet_name)
            last_col = max(len(sheet.columns), 1)
            current_row = 1

            if workbook_title:
                if last_col > 1:
                    worksheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=last_col)
                title_cell = worksheet.cell(current_row, 1, workbook_title)
                title_cell.font = title_font
                title_cell.alignment = centered
                current_row += 1

            description = str(sheet.description or "").strip()
            if description:
                if last_col > 1:
                    worksheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=last_col)
                description_cell = worksheet.cell(current_row, 1, description)
                description_cell.font = description_font
                current_row += 1

            header_row = current_row
            for col_idx, column in enumerate(sheet.columns, start=1):
                cell = worksheet.cell(header_row, col_idx, column.header)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
                cell.alignment = centered

            data_start_row = header_row + 1
            table_last_row = header_row
            if sheet.rows:
                for row_offset, row in enumerate(sheet.rows, start=data_start_row):
                    table_last_row = row_offset
                    for col_idx, column in enumerate(sheet.columns, start=1):
                        value = self._normalize_cell_value(row.get(column.key), column.type)
                        cell = worksheet.cell(row_offset, col_idx, value)
                        cell.border = border
                        cell.alignment = middle
                        cell.number_format = self._openpyxl_number_format(column.type)

            table_ref = f"A{header_row}:{get_column_letter(last_col)}{max(table_last_row, header_row)}"
            table = Table(displayName=self._table_name(index), ref=table_ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            worksheet.add_table(table)
            worksheet.auto_filter.ref = table_ref

            if sheet.freeze_header:
                worksheet.freeze_panes = f"A{header_row + 1}"

            for col_idx, column in enumerate(sheet.columns, start=1):
                letter = get_column_letter(col_idx)
                worksheet.column_dimensions[letter].width = self._calculate_column_width(column, sheet.rows)

            summary_row = table_last_row + 1
            if sheet.rows and sheet.summary and sheet.summary.formulas:
                label_cell = worksheet.cell(summary_row, 1, sheet.summary.label or "合计")
                label_cell.font = Font(bold=True)
                label_cell.fill = summary_fill
                label_cell.border = border

                for col_idx, column in enumerate(sheet.columns, start=1):
                    function_name = (sheet.summary.formulas or {}).get(column.key)
                    if not function_name:
                        continue

                    start_cell = f"{get_column_letter(col_idx)}{header_row + 1}"
                    end_cell = f"{get_column_letter(col_idx)}{table_last_row}"
                    summary_cell = worksheet.cell(
                        summary_row,
                        col_idx,
                        self._summary_formula(function_name, start_cell, end_cell),
                    )
                    summary_cell.font = Font(bold=True)
                    summary_cell.fill = summary_fill
                    summary_cell.border = border
                    summary_cell.number_format = self._openpyxl_number_format(column.type, summary=True)
                chart_anchor_row = summary_row + 2
            else:
                chart_anchor_row = table_last_row + 2

            if sheet.rows and sheet.chart and sheet.chart.series:
                category_col_idx = self._find_column_index(sheet.columns, sheet.chart.category_key)
                chart = self._build_openpyxl_chart(sheet.chart.type)
                if category_col_idx is not None and chart is not None:
                    categories = Reference(
                        worksheet,
                        min_col=category_col_idx + 1,
                        min_row=header_row + 1,
                        max_row=table_last_row,
                    )
                    for series in sheet.chart.series:
                        value_col_idx = self._find_column_index(sheet.columns, series.value_key)
                        if value_col_idx is None:
                            continue
                        values = Reference(
                            worksheet,
                            min_col=value_col_idx + 1,
                            min_row=header_row + 1,
                            max_row=table_last_row,
                        )
                        chart.add_data(values, titles_from_data=False)
                        if sheet.chart.type == "pie":
                            break

                    if chart.series:
                        chart.set_categories(categories)
                        if sheet.chart.title:
                            chart.title = sheet.chart.title
                        if hasattr(chart, "legend") and chart.legend:
                            chart.legend.position = "b"
                        if hasattr(chart, "x_axis") and sheet.chart.type != "pie":
                            chart.x_axis.title = sheet.columns[category_col_idx].header
                        anchor = f"{get_column_letter(last_col + 3)}{chart_anchor_row}"
                        worksheet.add_chart(chart, anchor)

            meta_row = max(chart_anchor_row, summary_row + 1 if sheet.rows else header_row + 2)
            meta_cell = worksheet.cell(meta_row, 1, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            meta_cell.font = meta_font

        workbook.save(file_path)

    def _build_formats(self, workbook):
        return {
            "title": workbook.add_format(
                {
                    "bold": True,
                    "font_size": 16,
                    "font_color": "#17324D",
                    "align": "center",
                    "valign": "vcenter",
                }
            ),
            "description": workbook.add_format(
                {
                    "font_size": 10,
                    "font_color": "#5F6B7A",
                    "italic": True,
                }
            ),
            "header": workbook.add_format(
                {
                    "bold": True,
                    "font_color": "#FFFFFF",
                    "bg_color": "#2F75B5",
                    "border": 1,
                    "align": "center",
                    "valign": "vcenter",
                }
            ),
            "text": workbook.add_format({"border": 1, "valign": "vcenter"}),
            "integer": workbook.add_format({"border": 1, "num_format": "#,##0", "valign": "vcenter"}),
            "number": workbook.add_format({"border": 1, "num_format": "#,##0.00", "valign": "vcenter"}),
            "currency": workbook.add_format({"border": 1, "num_format": '"¥"#,##0.00', "valign": "vcenter"}),
            "percentage": workbook.add_format({"border": 1, "num_format": "0.00%", "valign": "vcenter"}),
            "date": workbook.add_format({"border": 1, "num_format": "yyyy-mm-dd", "valign": "vcenter"}),
            "datetime": workbook.add_format({"border": 1, "num_format": "yyyy-mm-dd hh:mm", "valign": "vcenter"}),
            "summary_label": workbook.add_format(
                {
                    "bold": True,
                    "bg_color": "#D9E2F3",
                    "border": 1,
                }
            ),
            "summary_number": workbook.add_format(
                {
                    "bold": True,
                    "bg_color": "#D9E2F3",
                    "border": 1,
                    "num_format": "#,##0.00",
                }
            ),
            "summary_currency": workbook.add_format(
                {
                    "bold": True,
                    "bg_color": "#D9E2F3",
                    "border": 1,
                    "num_format": '"¥"#,##0.00',
                }
            ),
            "summary_percentage": workbook.add_format(
                {
                    "bold": True,
                    "bg_color": "#D9E2F3",
                    "border": 1,
                    "num_format": "0.00%",
                }
            ),
            "summary_integer": workbook.add_format(
                {
                    "bold": True,
                    "bg_color": "#D9E2F3",
                    "border": 1,
                    "num_format": "#,##0",
                }
            ),
            "table_body": workbook.add_format({"border": 1}),
            "meta": workbook.add_format({"font_size": 9, "font_color": "#7A7A7A"}),
        }

    def _normalize_cell_value(self, value: Any, cell_type: str) -> Any:
        if value is None:
            return ""

        if cell_type in {"integer", "number", "currency", "percentage"}:
            if isinstance(value, (int, float)):
                if cell_type == "percentage":
                    numeric_value = float(value)
                    return numeric_value / 100 if abs(numeric_value) > 1 else numeric_value
                return float(value) if cell_type in {"number", "currency"} else int(value)

            text = str(value).strip()
            if not text:
                return ""
            cleaned = text.replace(",", "").replace("¥", "").replace("￥", "")
            if cell_type == "percentage":
                has_percent_suffix = cleaned.endswith("%")
                numeric_value = float(cleaned[:-1].strip() if has_percent_suffix else cleaned)
                if has_percent_suffix or abs(numeric_value) > 1:
                    return numeric_value / 100
                return numeric_value
            if cell_type == "integer":
                return int(float(cleaned))
            return float(cleaned)

        if cell_type in {"date", "datetime"}:
            if isinstance(value, datetime):
                return value
            text = str(value).strip()
            if not text:
                return ""
            normalized = text.replace("Z", "+00:00")
            try:
                return datetime.fromisoformat(normalized)
            except ValueError:
                return text

        return value

    def _resolve_filename(self, ctx: TaskContext, planned_filename: Optional[str]) -> str:
        base_name = (ctx.output_filename or planned_filename or self._extract_filename(ctx.task_description) or "生成的工作簿.xlsx").strip()
        if not base_name.lower().endswith(".xlsx"):
            base_name += ".xlsx"
        return base_name

    def _coerce_workbook_spec(self, raw_spec: Any) -> ExcelWorkbookSpec:
        if isinstance(raw_spec, ExcelWorkbookSpec):
            return raw_spec

        normalized = copy.deepcopy(raw_spec or {})
        if not isinstance(normalized, dict):
            raise ValueError("workbook spec 必须是对象")

        if "filename" not in normalized and normalized.get("file_name"):
            normalized["filename"] = normalized.get("file_name")

        for sheet in normalized.get("sheets", []) or []:
            formulas: dict[str, str] = {}
            columns = sheet.get("columns") or []
            for column in columns:
                if not isinstance(column, dict):
                    continue
                raw_type = str(column.get("type", "text")).strip().lower()
                if raw_type == "percent":
                    column["type"] = "percentage"
                summary_name = column.pop("summary", None)
                if summary_name:
                    formulas[str(column.get("key"))] = str(summary_name)

            rows = sheet.get("rows") or []
            if rows and isinstance(rows[0], list):
                keys = [str(column.get("key")) for column in columns]
                sheet["rows"] = [
                    {keys[idx]: row[idx] if idx < len(row) else None for idx in range(len(keys))}
                    for row in rows
                ]

            chart = sheet.get("chart")
            if isinstance(chart, dict) and chart.get("value_key") and not chart.get("series"):
                chart["series"] = [
                    {
                        "name": chart.get("title") or chart.get("value_key"),
                        "value_key": chart.get("value_key"),
                    }
                ]

            if formulas and not sheet.get("summary"):
                sheet["summary"] = {
                    "label": "合计",
                    "formulas": formulas,
                }

        return ExcelWorkbookSpec.model_validate(normalized)

    def _extract_filename(self, task_description: str) -> str:
        match = re.search(r"《(.+?)》", task_description)
        if match:
            return match.group(1)

        match = re.search(r"生成(.+?)(报表|台账|明细|清单|工作簿|表格|分析)", task_description)
        if match:
            return match.group(1) + match.group(2)

        return ""

    def _build_creation_context(self, ctx: TaskContext) -> str:
        base_context = build_query_focused_context(
            query=ctx.task_description,
            memory_dfs=ctx.memory_dfs,
            context_title="筛选后的可用数据（只保留与 Excel 任务直接相关的数据）",
            max_sources=3,
            max_rows=24,
            max_chars=18000,
            max_chars_per_source=5000,
        )
        result_text = extract_query_focused_result_text(
            query=ctx.task_description,
            execution_results=ctx.execution_results,
            max_blocks=8,
            max_chars=8000,
        )
        if not result_text:
            return base_context

        full_block = (
            "\n\n## 筛选后的检索片段（只保留与 Excel 任务直接相关的内容）\n"
            f"{result_text}\n"
        )
        return f"{base_context}{full_block}" if base_context else full_block

    def _has_meaningful_rows(self, workbook_spec: ExcelWorkbookSpec) -> bool:
        for sheet in workbook_spec.sheets:
            if not sheet.columns:
                continue
            for row in sheet.rows:
                if any(self._is_meaningful_value(row.get(column.key)) for column in sheet.columns):
                    return True
        return False

    def _is_meaningful_value(self, value: Any) -> bool:
        if value is None:
            return False
        if isinstance(value, str):
            return bool(value.strip())
        return True

    def _calculate_column_width(self, column: ExcelColumnSpec, rows: list[dict[str, Any]]) -> float:
        max_len = len(str(column.header or ""))
        for row in rows[:200]:
            value = row.get(column.key)
            if value is None:
                continue
            display = str(value)
            max_len = max(max_len, len(display))
        return min(max(max_len + 2, 10), 28)

    def _find_column_index(self, columns: list[ExcelColumnSpec], key: str) -> Optional[int]:
        for index, column in enumerate(columns):
            if column.key == key:
                return index
        return None

    def _sanitize_sheet_name(self, raw_name: str, index: int) -> str:
        sanitized = re.sub(r"[\[\]:*?/\\]", "_", str(raw_name or "").strip()) or f"Sheet{index}"
        return sanitized[:31]

    def _chart_type(self, chart_type: str) -> str:
        return chart_type if chart_type in {"column", "bar", "line", "pie", "area", "scatter"} else "column"

    def _value_format_name(self, cell_type: str) -> str:
        mapping = {
            "integer": "integer",
            "number": "number",
            "currency": "currency",
            "percentage": "percentage",
            "date": "date",
            "datetime": "datetime",
        }
        return mapping.get(cell_type, "text")

    def _summary_format_name(self, cell_type: str) -> str:
        mapping = {
            "integer": "summary_integer",
            "currency": "summary_currency",
            "percentage": "summary_percentage",
        }
        return mapping.get(cell_type, "summary_number")

    def _summary_formula(self, function_name: str, start_cell: str, end_cell: str) -> str:
        subtotal_codes = {
            "sum": 109,
            "average": 101,
            "count": 103,
            "min": 105,
            "max": 104,
        }
        code = subtotal_codes.get(str(function_name or "").strip().lower(), 109)
        return f"=SUBTOTAL({code},{start_cell}:{end_cell})"

    def _openpyxl_number_format(self, cell_type: str, summary: bool = False) -> str:
        if cell_type == "integer":
            return "#,##0"
        if cell_type == "number":
            return "#,##0.00"
        if cell_type == "currency":
            return '"¥"#,##0.00'
        if cell_type == "percentage":
            return "0.00%"
        if cell_type == "date":
            return "yyyy-mm-dd"
        if cell_type == "datetime":
            return "yyyy-mm-dd hh:mm"
        return "General"

    def _build_openpyxl_chart(self, chart_type: str):
        from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, ScatterChart

        normalized = self._chart_type(chart_type)
        if normalized == "line":
            chart = LineChart()
        elif normalized == "pie":
            chart = PieChart()
        elif normalized == "area":
            chart = AreaChart()
        elif normalized == "scatter":
            chart = ScatterChart()
        else:
            chart = BarChart()
            chart.type = "bar" if normalized == "bar" else "col"

        chart.style = 10
        chart.height = 7
        chart.width = 12
        return chart

    def _table_name(self, index: int) -> str:
        return f"Table{index}"

    def _error_result(self, ctx: TaskContext, message: str, reason_code: str, retryable: bool) -> Dict[str, Any]:
        return {
            "success": False,
            "error": message,
            "quality_signal": {
                "verdict": "fail",
                "reason_code": reason_code,
                "confidence": 0.9,
                "retryable": retryable,
            },
            "meta": {
                "worker_round": ctx.round_index,
                "stop_reason": reason_code,
            },
        }
